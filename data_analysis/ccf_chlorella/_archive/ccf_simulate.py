# -*- coding: utf-8 -*-
"""
CCF（面心中心复合设计）—— 小球藻微反应器培养：模拟数据生成 + RSM 分析 + 讲解报告
================================================================================
因子（3 个）：
  x1 光照强度   30–210 μmol/m²/s（中心 120，半幅 90）
  x2 氨氮浓度   20–300 mg/L      （中心 160，半幅 140）
  x3 无机碳浓度 75–565 mg/L      （中心 320，半幅 245）

设计：CCF（α=1），N = 2^3(析因) + 2×3(轴点) + 6(中心点) = 20 组
响应：μmax —— 小球藻最大比生长速率（d⁻¹）（第三轮需求：由生物量 Y 改为 μmax）
真实模型驻点（编码）：x1*=0.7, x2*=0.9, x3*=0.6（用户指定）
真实交互作用：β12>0（光照×氨氮 协同）、β13>0（光照×无机碳 协同）、β23<0（氨氮×无机碳 拮抗）
范围锁定：数据先在基准标度上生成（驻点与显著性已验证），再经仿射标定 y'=a·y+b
        精确锁定 μmax ∈ [0.0571, 0.2072]；仿射变换不改变驻点与 t/p、R²、失拟统计量

输出（均写入 E:/pythonProject/Microalgae_Identification_YOLOv11/data_analysis/ccf_chlorella/）：
  CCF_小球藻_20组模拟数据.xlsx      —— 4 个工作表的数据/分析文件
  figures/fig1..fig6.png           —— 6 张分析图
  widget_interaction.svg            —— 内联讲解用交互作用图（SVG）
  CCF_小球藻_讲解报告.html          —— 自包含讲解报告（图片 base64 内嵌）
  summary.json                      —— 全部关键数字（供核对）
"""
import os
import json
import base64
import itertools
from datetime import datetime

import numpy as np
from scipy import stats as st

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# 0. 路径与全局配置
# ----------------------------------------------------------------------------
BASE = "E:/pythonProject/Microalgae_Identification_YOLOv11/data_analysis/ccf_chlorella"
FIG_DIR = os.path.join(BASE, "figures")
os.makedirs(FIG_DIR, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DengXian", "sans-serif"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.dpi"] = 110

# 因子定义：名称 / 单位 / 低 / 高
FACTORS = [
    dict(key="x1", name="光照强度", en="Light intensity", unit="μmol/m²/s", lo=30,  hi=210),
    dict(key="x2", name="氨氮浓度", en="Ammonia nitrogen", unit="mg/L",     lo=20,  hi=300),
    dict(key="x3", name="无机碳浓度", en="Inorganic carbon", unit="mg/L",   lo=75,  hi=565),
]
CENTER = [(f["lo"] + f["hi"]) / 2 for f in FACTORS]      # 120, 160, 320
HALF   = [(f["hi"] - f["lo"]) / 2 for f in FACTORS]      # 90, 140, 245

def to_actual(i, x):
    """编码值 -> 实际值"""
    return CENTER[i] + HALF[i] * x

# 目标最优编码值（用户指定）
TARGET = np.array([0.7, 0.9, 0.6])

# 响应定义与范围锁定（第三轮：μmax 取代生物量 Y）
RESP_NAME = "μmax"                 # 最大比生长速率
RESP_UNIT = "d⁻¹"
TARGET_MIN = 0.0571                # 20 组数据中 μmax 的最小值（精确锁定）
TARGET_MAX = 0.2072                # 20 组数据中 μmax 的最大值（精确锁定）
A_SCL, B_SCL = 0.0, 0.0            # 仿射标定系数（main() 中依数据计算）

# ----------------------------------------------------------------------------
# 1. CCF 设计矩阵（标准序）
# ----------------------------------------------------------------------------
def build_design():
    pts, typ = [], []
    for c in itertools.product([-1, 1], repeat=3):        # 8 个析因点
        pts.append(c); typ.append("析因点")
    for i in range(3):                                     # 6 个轴点（α=1 面心）
        for s in (-1, 1):
            p = [0, 0, 0]; p[i] = s
            pts.append(tuple(p)); typ.append("轴点")
    for _ in range(6):                                     # 6 个中心点
        pts.append((0, 0, 0)); typ.append("中心点")
    return np.array(pts, dtype=float), typ

# ----------------------------------------------------------------------------
# 2. “真实”二次模型（构造数据用，基准标度）
#    驻点被精确设计在 (0.7, 0.9, 0.6)；Hessian 负定 → 该驻点为极大值
#    μmax 标度下的真实系数 = 仿射换算（截距 B_SCL+A_SCL·B0，其余 A_SCL·B*）
# ----------------------------------------------------------------------------
B0 = 2.26
B_LIN  = (0.168, 0.594, 0.315)          # β1  β2  β3
B_INT  = (0.18,  0.15, -0.12)           # β12 β13 β23
B_QUAD = (-0.30, -0.36, -0.26)          # β11 β22 β33

def true_model(x1, x2, x3):
    return (B0
            + B_LIN[0]*x1 + B_LIN[1]*x2 + B_LIN[2]*x3
            + B_INT[0]*x1*x2 + B_INT[1]*x1*x3 + B_INT[2]*x2*x3
            + B_QUAD[0]*x1*x1 + B_QUAD[1]*x2*x2 + B_QUAD[2]*x3*x3)

def verify_true_stationary():
    """校验真实模型驻点 == TARGET 且为极大值"""
    A = np.array([[2*B_QUAD[0], B_INT[0],     B_INT[1]],
                  [B_INT[0],     2*B_QUAD[1], B_INT[2]],
                  [B_INT[1],     B_INT[2],    2*B_QUAD[2]]])
    b = -np.array(B_LIN)
    xs = np.linalg.solve(A, b)
    ev = np.linalg.eigvalsh(A)
    assert np.allclose(xs, TARGET, atol=1e-12), f"真实模型驻点 {xs} != {TARGET}"
    assert np.all(ev < 0), "Hessian 非负定，驻点不是极大值"
    return xs

# ----------------------------------------------------------------------------
# 3. 最小二乘拟合 + 方差分析
# ----------------------------------------------------------------------------
TERM_NAMES = ["截距", "x1 光照", "x2 氨氮", "x3 无机碳",
              "x1x2 光照×氨氮", "x1x3 光照×无机碳", "x2x3 氨氮×无机碳",
              "x1² 光照²", "x2² 氨氮²", "x3² 无机碳²"]

def X_matrix(pts):
    cols = []
    for x1, x2, x3 in pts:
        cols.append([1, x1, x2, x3, x1*x2, x1*x3, x2*x3, x1*x1, x2*x2, x3*x3])
    return np.array(cols)

def model_at(beta, x1, x2, x3):
    Xr = X_matrix([(x1, x2, x3)])
    return float((Xr @ beta)[0])

def analyze(pts, y):
    """完整 RSM 分析：系数、t/p、ANOVA、失拟、R2、驻点、预测区间"""
    X = X_matrix(pts)
    n, p = X.shape                      # 20, 10
    XtXi = np.linalg.inv(X.T @ X)
    beta = XtXi @ X.T @ y
    resid = y - X @ beta
    RSS = float(resid @ resid)
    dfe = n - p                         # 10
    s2 = RSS / dfe
    s = np.sqrt(s2)
    se = np.sqrt(np.diag(s2 * XtXi))
    tvals = beta / se
    pvals = 2 * (1 - st.t.cdf(np.abs(tvals), dfe))

    # 纯误差（6 个中心点）与失拟
    yc = y[14:20]
    PE = float(((yc - yc.mean())**2).sum()); dfpe = len(yc) - 1
    LF = RSS - PE; dflf = dfe - dfpe
    F_lof = (LF/dflf) / (PE/dfpe) if PE > 0 else np.inf
    p_lof = 1 - st.f.cdf(F_lof, dflf, dfpe)

    # 总体回归
    ybar = y.mean()
    TSS = float(((y - ybar)**2).sum())
    RegSS = TSS - RSS
    F_reg = (RegSS/(p-1)) / s2
    p_reg = 1 - st.f.cdf(F_reg, p-1, dfe)
    R2 = 1 - RSS/TSS
    adjR2 = 1 - (1-R2)*(n-1)/dfe

    # 驻点
    A = np.array([[2*beta[7], beta[4],  beta[5]],
                  [beta[4],   2*beta[8], beta[6]],
                  [beta[5],   beta[6],  2*beta[9]]])
    xs = np.linalg.solve(A, -np.array(beta[1:4]))
    ev = np.linalg.eigvalsh(A)
    yopt = model_at(beta, *xs)

    # 最优点处的置信/预测区间
    x0 = X_matrix([tuple(xs)])[0]
    lev = float(x0 @ XtXi @ x0)
    tcrit = st.t.ppf(0.975, dfe)
    ci_half = tcrit * s * np.sqrt(lev)
    pi_half = tcrit * s * np.sqrt(1 + lev)

    return dict(beta=beta, se=se, t=tvals, p=pvals, s=s, RSS=RSS, dfe=dfe,
                PE=PE, dfpe=dfpe, LF=LF, dflf=dflf, F_lof=F_lof, p_lof=p_lof,
                RegSS=RegSS, TSS=TSS, F_reg=F_reg, p_reg=p_reg,
                R2=R2, adjR2=adjR2, xs=xs, eig=ev, yopt=yopt,
                ci_half=ci_half, pi_half=pi_half, tcrit=float(tcrit),
                center_mean=float(yc.mean()), center_sd=float(yc.std(ddof=1)))

# ----------------------------------------------------------------------------
# 4. 种子搜索：找到满足全部统计与构造要求的模拟数据（基准标度）
# ----------------------------------------------------------------------------
def seed_search():
    candidates = []
    for sigma in (0.08, 0.07, 0.06):
        acc = 0
        for seed in range(1, 1500):
            rng = np.random.default_rng(seed)
            y = np.array([true_model(*p) for p in PTS]) + rng.normal(0, sigma, 20)
            y = np.round(y, 3)
            if y.min() <= 0.05:
                continue
            if not (0.015 <= y[14:20].std(ddof=1) <= 0.12):
                continue
            r = analyze(PTS, y)
            # 显著性
            if max(r["p"][4], r["p"][5], r["p"][6]) > 0.01:   # 交互项 p<0.01
                continue
            if max(r["p"][7], r["p"][8], r["p"][9]) > 0.01:   # 二次项 p<0.01
                continue
            if max(r["p"][1], r["p"][2], r["p"][3]) > 0.05:   # 线性项 p<0.05
                continue
            # 系数符号符合机理设定
            if not (r["beta"][4] > 0 and r["beta"][5] > 0 and r["beta"][6] < 0
                    and r["beta"][7] < 0 and r["beta"][8] < 0 and r["beta"][9] < 0):
                continue
            if r["p_lof"] <= 0.10:                            # 失拟不显著
                continue
            if not (0.93 <= r["R2"] <= 0.998):
                continue
            # 驻点：贴近目标且在立方体内、且为极大值
            if r["eig"].max() >= 0:
                continue
            dev = np.abs(r["xs"] - TARGET)
            if dev.max() > 0.075 or np.abs(r["xs"]).max() > 0.95:
                continue
            acc += 1
            candidates.append((float(dev.max()), -sigma, seed, y.copy(), r))
        print(f"  sigma={sigma:.2f}: 通过筛选 {acc} 个候选")
    if not candidates:
        raise RuntimeError("未找到满足要求的种子，请放宽条件")
    candidates.sort(key=lambda c: (round(c[0], 4), c[1]))
    best = candidates[0]
    print(f"  选定: seed={best[2]}, sigma={-best[1]:.2f}, 驻点最大偏差={best[0]:.4f}")
    return best[2], -best[1], best[3]

# ----------------------------------------------------------------------------
# 5. 图表
# ----------------------------------------------------------------------------
C_LIN, C_INT, C_QUAD, C_CEN = "#3b82f6", "#f59e0b", "#10b981", "#8b5cf6"

def ticks_actual(ax, which, i):
    f = FACTORS[i]
    vals = [f["lo"], int(CENTER[i]), f["hi"]]
    if which == "x":
        ax.set_xticks([-1, 0, 1]); ax.set_xticklabels([str(v) for v in vals])
    else:
        ax.set_yticks([-1, 0, 1]); ax.set_yticklabels([str(v) for v in vals])

def fig1_design():
    fig = plt.figure(figsize=(8.6, 6.8))
    ax = fig.add_subplot(111, projection="3d")
    # 立方体棱
    corners = list(itertools.product([-1, 1], repeat=3))
    for a, b in itertools.combinations(corners, 2):
        d = sum(abs(u-v) for u, v in zip(a, b))
        if abs(d - 2) < 1e-9:
            ax.plot(*zip(a, b), color="#94a3b8", lw=1.2, alpha=0.8)
    P = PTS
    m = [t == "析因点" for t in TYPES]
    ax.scatter(*P[m].T, c="#f59e0b", s=80, edgecolors="#92400e", lw=0.8,
               depthshade=False, label="析因点 2³ = 8（±1 组合）")
    m = [t == "轴点" for t in TYPES]
    ax.scatter(*P[m].T, c="#3b82f6", s=80, edgecolors="#1e3a8a", lw=0.8, marker="s",
               depthshade=False, label="轴点 2k = 6（α=1，面心）")
    ax.scatter([0], [0], [0], c="#10b981", s=300, marker="*", edgecolors="#065f46",
               lw=0.8, depthshade=False, label="中心点 ×6（0,0,0）")
    ax.set_xlabel(f"光照强度 (μmol/m²/s)\n[-1]=30  [0]=120  [+1]=210", labelpad=8)
    ax.set_ylabel(f"氨氮浓度 (mg/L)\n[-1]=20  [0]=160  [+1]=300", labelpad=8)
    ax.set_zlabel("无机碳 (mg/L)", labelpad=6)
    ax.set_xticks([-1, 0, 1]); ax.set_yticks([-1, 0, 1]); ax.set_zticks([-1, 0, 1])
    ax.set_title("CCF 设计空间：20 组 = 8 析因点 + 6 轴点(α=1) + 6 中心点",
                 fontsize=12, pad=2)
    ax.view_init(elev=16, azim=-58)
    ax.legend(loc="upper left", bbox_to_anchor=(0.02, 0.98), fontsize=9)
    p = os.path.join(FIG_DIR, "fig1_design.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

def _curve(beta, free_i, line_i, line_v, fixed_i, fixed_v, xs=np.linspace(-1, 1, 41)):
    """第三因子固定时，某因子扫描、另一因子取给定水平的响应曲线"""
    out = []
    for x in xs:
        v = [None, None, None]
        v[free_i] = x; v[line_i] = line_v; v[fixed_i] = fixed_v
        out.append(model_at(beta, *v))
    return xs, np.array(out)

def fig2_interaction(beta):
    panels = [
        dict(free=0, line=1, fixed=2, title="(a) 光照 × 氨氮（无机碳 = 320 mg/L）",
             xl="光照强度（编码）", lines="氨氮", li=1),
        dict(free=0, line=2, fixed=1, title="(b) 光照 × 无机碳（氨氮 = 160 mg/L）",
             xl="光照强度（编码）", lines="无机碳", li=2),
        dict(free=1, line=2, fixed=0, title="(c) 氨氮 × 无机碳（光照 = 120 μmol/m²/s）",
             xl="氨氮浓度（编码）", lines="无机碳", li=2),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(13.2, 4.4), sharey=False)
    for ax, pn in zip(axes, panels):
        for lv, c, ls, mk in [(-1, "#2563eb", "-", "o"), (0, "#64748b", "--", "s"),
                              (1, "#dc2626", "-", "D")]:
            xs_, ys_ = _curve(beta, pn["free"], pn["line"], lv, pn["fixed"], 0)
            act = to_actual(pn["li"], lv)
            ax.plot(xs_, ys_, color=c, ls=ls, lw=2, marker=mk, ms=4,
                    label=f"{pn['lines']} {lv:+d}（{act:g} {FACTORS[pn['li']]['unit']}）")
        ax.set_title(pn["title"], fontsize=11)
        ax.set_xlabel(pn["xl"])
        ax.grid(alpha=0.25, ls=":")
        ax.legend(fontsize=8.5, loc="best")
    axes[0].set_ylabel("μmax (d$^{-1}$)")
    fig.suptitle("交互作用图（基于拟合模型；折线不平行即存在交互作用）", fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    p = os.path.join(FIG_DIR, "fig2_interaction.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

def fig3_contour(beta, xs):
    panels = [
        dict(i=0, j=1, k=2, title="(a) 光照–氨氮 平面（无机碳固定于最优编码 %.2f）" % xs[2]),
        dict(i=0, j=2, k=1, title="(b) 光照–无机碳 平面（氨氮固定于最优编码 %.2f）" % xs[1]),
        dict(i=1, j=2, k=0, title="(c) 氨氮–无机碳 平面（光照固定于最优编码 %.2f）" % xs[0]),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(13.6, 4.6))
    g = np.linspace(-1, 1, 200)
    for ax, pn in zip(axes, panels):
        GX, GY = np.meshgrid(g, g)
        V = [None, None, None]; V[pn["k"]] = xs[pn["k"]]
        Z = np.zeros_like(GX)
        for a in range(GX.shape[0]):
            for b in range(GX.shape[1]):
                V[pn["i"]] = GX[a, b]; V[pn["j"]] = GY[a, b]
                Z[a, b] = model_at(beta, *V)
        cf = ax.contourf(GX, GY, Z, levels=16, cmap="RdYlGn")
        cs = ax.contour(GX, GY, Z, levels=16, colors="k", linewidths=0.35, alpha=0.45)
        ax.clabel(cs, fmt="%.3f", fontsize=6.5)
        ax.plot(xs[pn["i"]], xs[pn["j"]], marker="*", ms=17, color="#0f172a",
                mec="white", mew=0.8, ls="none", zorder=5)
        ax.axvline(xs[pn["i"]], color="#0f172a", lw=0.8, ls="--", alpha=0.55)
        ax.axhline(xs[pn["j"]], color="#0f172a", lw=0.8, ls="--", alpha=0.55)
        ax.set_title(pn["title"], fontsize=10.5)
        ax.set_xlabel(f"{FACTORS[pn['i']]['name']}（{FACTORS[pn['i']]['unit']}）")
        ax.set_ylabel(f"{FACTORS[pn['j']]['name']}（{FACTORS[pn['j']]['unit']}）")
        ticks_actual(ax, "x", pn["i"]); ticks_actual(ax, "y", pn["j"])
        fig.colorbar(cf, ax=ax, shrink=0.9, label="μmax (d$^{-1}$)")
    fig.suptitle("等高线图（★ = 最优点；黑虚线十字定位最优编码水平）", fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.93])
    p = os.path.join(FIG_DIR, "fig3_contour.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

def fig4_surface(beta, xs):
    fig = plt.figure(figsize=(9.2, 6.6))
    ax = fig.add_subplot(111, projection="3d")
    g = np.linspace(-1, 1, 60)
    GX, GY = np.meshgrid(g, g)
    Z = np.zeros_like(GX)
    for a in range(GX.shape[0]):
        for b in range(GX.shape[1]):
            Z[a, b] = model_at(beta, GX[a, b], GY[a, b], xs[2])
    ax.plot_surface(GX, GY, Z, cmap="viridis", alpha=0.92, rstride=1, cstride=1,
                    linewidth=0, antialiased=True)
    ax.contour(GX, GY, Z, zdir="z", offset=Z.min()-0.015, levels=12, cmap="viridis",
               linewidths=0.8)
    yopt = model_at(beta, xs[0], xs[1], xs[2])
    ax.scatter([xs[0]], [xs[1]], [yopt], marker="*", s=280, color="#f43f5e",
               edgecolors="white", lw=1.0)
    ax.text(xs[0], xs[1], yopt + 0.006,
            "最优点 (%.2f, %.2f)\nμmax = %.4f" % (xs[0], xs[1], yopt),
            fontsize=10, ha="center", fontweight="bold")
    ax.set_xlabel("光照强度（编码）")
    ax.set_ylabel("氨氮浓度（编码）")
    ax.set_zlabel("μmax (d$^{-1}$)")
    ax.set_title("响应面三维图（无机碳固定于最优编码 %.2f，即 %d mg/L）"
                 % (xs[2], to_actual(2, xs[2])), fontsize=12)
    ax.view_init(elev=24, azim=-55)
    p = os.path.join(FIG_DIR, "fig4_surface.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

def fig5_pareto(r):
    idx = list(range(1, 10))
    idx.sort(key=lambda i: abs(r["t"][i]))
    labels = [TERM_NAMES[i] for i in idx]
    tv = [abs(r["t"][i]) for i in idx]
    cols = [C_LIN if i <= 3 else (C_INT if i <= 6 else C_QUAD) for i in idx]
    fig, ax = plt.subplots(figsize=(9.4, 4.6))
    bars = ax.barh(range(len(idx)), tv, color=cols, edgecolor="white", height=0.68)
    for k, (bar, v) in enumerate(zip(bars, tv)):
        ax.text(v + 0.12, bar.get_y() + bar.get_height()/2, f"{v:.1f}",
                va="center", fontsize=9)
    tcrit = st.t.ppf(0.975, r["dfe"])
    ax.axvline(tcrit, color="#dc2626", ls="--", lw=1.4)
    ax.text(tcrit + 0.1, -0.62, f"$t_{{0.05}}$({r['dfe']}) = {tcrit:.2f}",
            color="#dc2626", fontsize=9.5)
    ax.set_yticks(range(len(idx))); ax.set_yticklabels(labels, fontsize=9.5)
    ax.set_xlabel("标准化效应 |t|")
    ax.set_title("效应 Pareto 图（越靠右越显著；全部越过显著性门槛）", fontsize=12)
    handles = [Line2D([], [], marker="s", ls="", color=C_LIN, label="线性项"),
               Line2D([], [], marker="s", ls="", color=C_INT, label="交互项"),
               Line2D([], [], marker="s", ls="", color=C_QUAD, label="二次项")]
    ax.legend(handles=handles, loc="lower right", fontsize=9)
    ax.set_xlim(0, max(tv) * 1.14)
    p = os.path.join(FIG_DIR, "fig5_pareto.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

def fig6_obspred(beta, r):
    yhat = X_matrix(PTS) @ beta
    fig, ax = plt.subplots(figsize=(6.4, 5.6))
    for typ, c, mk in [("析因点", "#f59e0b", "o"), ("轴点", "#3b82f6", "s"),
                       ("中心点", "#10b981", "D")]:
        m = [t == typ for t in TYPES]
        ax.scatter(yhat[m], Y_DATA[m], c=c, marker=mk, s=64, alpha=0.9,
                   edgecolors="white", lw=0.8, label=typ)
    lims = [min(yhat.min(), Y_DATA.min()) - 0.008, max(yhat.max(), Y_DATA.max()) + 0.008]
    ax.plot(lims, lims, color="#64748b", ls="--", lw=1.2)
    ax.set_xlabel("模型预测值 ŷ (d$^{-1}$)")
    ax.set_ylabel("模拟观测值 μmax (d$^{-1}$)")
    ax.set_title(f"观测值 vs 预测值（R² = {r['R2']:.4f}）", fontsize=12)
    ax.grid(alpha=0.25, ls=":")
    ax.legend(fontsize=9)
    ax.set_aspect("equal", adjustable="box")
    p = os.path.join(FIG_DIR, "fig6_obspred.png")
    fig.savefig(p, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p

# ----------------------------------------------------------------------------
# 6. 内联讲解用 SVG：三联交互作用图（μmax 标度）
# ----------------------------------------------------------------------------
def svg_interaction(beta):
    """内联讲解用 SVG：三联交互作用图（μmax 标度；配色与排版遵循浅色主题设计系统）"""
    W, H = 680, 330
    panels = [
        dict(free=0, line=1, fixed=2, title="光照 × 氨氮",
             note="无机碳固定 320 mg/L", xl="光照强度（编码）"),
        dict(free=0, line=2, fixed=1, title="光照 × 无机碳",
             note="氨氮固定 160 mg/L", xl="光照强度（编码）"),
        dict(free=1, line=2, fixed=0, title="氨氮 × 无机碳",
             note="光照固定 120 μmol/m²/s", xl="氨氮浓度（编码）"),
    ]
    STYLE = {1: ("#378ADD", "none"), 0: ("#888780", "5 4"), -1: ("#EF9F27", "none")}
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" role="img" '
        f'font-family="Microsoft YaHei, PingFang SC, sans-serif">',
        '<title>三对因子的交互作用图（μmax 标度）</title>',
        '<desc>三个面板分别展示光照×氨氮、光照×无机碳、氨氮×无机碳的交互作用：横轴为一个因子的'
        '编码水平（-1 到 +1），三条折线为另一因子取低、中、高水平时的 μmax 响应曲线。'
        '折线不平行表明一个因子的效应大小随另一因子的水平而改变，即存在显著交互作用。</desc>',
    ]
    ytop, ybot = 78, 210
    for k, pn in enumerate(panels):
        base = 76 + k * 200
        px0, px1 = base + 8, base + 144
        curves = {}
        for lv in (-1, 0, 1):
            xs_, ys_ = _curve(beta, pn["free"], pn["line"], lv, pn["fixed"], 0)
            curves[lv] = (xs_, ys_)
        allv = np.concatenate([curves[lv][1] for lv in (-1, 0, 1)])
        ymin, ymax = float(allv.min()), float(allv.max())
        padv = (ymax - ymin) * 0.10 + 1e-6
        ymin, ymax = ymin - padv, ymax + padv

        def SX(x): return px0 + (x + 1) / 2 * (px1 - px0)
        def SY(y): return ybot - (y - ymin) / (ymax - ymin) * (ybot - ytop)

        parts.append(f'<rect x="{px0-6}" y="{ytop-6}" width="{px1-px0+12}" '
                     f'height="{ybot-ytop+12}" fill="#FAFAF8" stroke="#D3D1C7" '
                     f'stroke-width="1" rx="6"/>')
        parts.append(f'<text x="{(px0+px1)/2}" y="54" text-anchor="middle" '
                     f'font-size="12.5" font-weight="500" fill="#2C2C2A">{pn["title"]}</text>')
        parts.append(f'<text x="{(px0+px1)/2}" y="69" text-anchor="middle" '
                     f'font-size="11" fill="#5F5E5A">{pn["note"]}</text>')
        for xv in (-1, 0, 1):
            parts.append(f'<line x1="{SX(xv):.1f}" y1="{ytop}" x2="{SX(xv):.1f}" '
                         f'y2="{ybot}" stroke="#D3D1C7" stroke-width="1"/>')
        for yv in np.linspace(ymin, ymax, 3)[1:-1]:
            parts.append(f'<line x1="{px0}" y1="{SY(yv):.1f}" x2="{px1}" '
                         f'y2="{SY(yv):.1f}" stroke="#D3D1C7" stroke-width="1"/>')
        parts.append(f'<line x1="{px0}" y1="{ybot}" x2="{px1}" y2="{ybot}" '
                     f'stroke="#888780" stroke-width="1.25"/>')
        parts.append(f'<line x1="{px0}" y1="{ytop}" x2="{px0}" y2="{ybot}" '
                     f'stroke="#888780" stroke-width="1.25"/>')
        for xv, xlab in ((-1, "-1"), (0, "0"), (1, "+1")):
            parts.append(f'<text x="{SX(xv):.1f}" y="{ybot+15}" text-anchor="middle" '
                         f'font-size="11" fill="#5F5E5A">{xlab}</text>')
        for yv in np.linspace(ymin, ymax, 3):
            parts.append(f'<text x="{px0-4}" y="{SY(yv)+3.5:.1f}" text-anchor="end" '
                         f'font-size="11" fill="#5F5E5A">{yv:.3f}</text>')
        for lv in (-1, 0, 1):
            c, dash = STYLE[lv]
            xs_, ys_ = curves[lv]
            pts_ = " ".join(f"{SX(a):.1f},{SY(b):.1f}" for a, b in zip(xs_, ys_))
            parts.append(f'<polyline points="{pts_}" fill="none" stroke="{c}" '
                         f'stroke-width="2.25" stroke-dasharray="{dash}" '
                         f'stroke-linejoin="round" stroke-linecap="round"/>')
        parts.append(f'<text x="{(px0+px1)/2}" y="240" text-anchor="middle" '
                     f'font-size="11.5" fill="#5F5E5A">{pn["xl"]}</text>')
    legend = [("高水平 (+1)", "#378ADD", "none"), ("中心 (0)", "#888780", "5 4"),
              ("低水平 (−1)", "#EF9F27", "none")]
    lx = 140
    for name, c, dash in legend:
        parts.append(f'<line x1="{lx}" y1="259" x2="{lx+24}" y2="259" stroke="{c}" '
                     f'stroke-width="2.5" stroke-dasharray="{dash}" '
                     f'stroke-linecap="round"/>')
        parts.append(f'<text x="{lx+30}" y="263" font-size="11.5" fill="#5F5E5A">{name}</text>')
        lx += 130
    parts.append(f'<text x="{W/2}" y="288" text-anchor="middle" font-size="11.5" '
                 f'fill="#5F5E5A">纵轴为 μmax（d⁻¹）。折线不平行 —— 一个因子的效应大小'
                 f'随另一因子的水平而改变，即存在显著交互作用。</text>')
    parts.append("</svg>")
    svg = "".join(parts)
    p = os.path.join(BASE, "widget_interaction.svg")
    with open(p, "w", encoding="utf-8") as f:
        f.write(svg)
    return p

# ----------------------------------------------------------------------------
# 7. Excel 工作簿
# ----------------------------------------------------------------------------
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10.5)
BODY_FONT = Font(name="微软雅黑", size=10.5)
TYPE_FILL = {"析因点": PatternFill("solid", fgColor="FEF3C7"),
             "轴点": PatternFill("solid", fgColor="DBEAFE"),
             "中心点": PatternFill("solid", fgColor="D1FAE5")}

def style_header(ws, row, ncol, start=1):
    for c in range(start, start + ncol):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def make_xlsx(r, seed, sigma, run_no):
    wb = openpyxl.Workbook()
    # --- Sheet 1: 说明 ---
    ws = wb.active; ws.title = "说明"
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 88
    rows = [
        ("CCF（面心中心复合）设计 —— 小球藻微反应器培养模拟数据（响应：μmax）", ""),
        ("文件性质", "教学演示用模拟数据集（由已知“真实”二次模型 + 正态随机误差生成），非实测数据"),
        ("响应指标 μmax", "小球藻最大比生长速率（单位 d⁻¹），由批式培养生长曲线拟合（如 Logistic 模型）取得"),
        ("数据范围约束", "20 组 μmax 精确锁定于 [0.0571, 0.2072] d⁻¹（最小值 0.0571、最大值 0.2072，经仿射标定）"),
        ("设计类型", "Central Composite Face-centered（CCF，α = 1），3 因子，20 组"),
        ("设计构成", "2³ = 8 个析因点 + 2×3 = 6 个轴点（面心）+ 6 个中心点"),
        ("随机化", "表中“运行序”为随机化执行顺序（可复现，种子见下）"),
        ("模拟参数", f"构造噪声 σ ≈ {A_SCL*sigma:.4f} d⁻¹（正态；基准标度 σ = {sigma:.2f} 经仿射"
                     f"标度 a = {A_SCL:.6f}、b = {B_SCL:.6f} 换算），随机种子 seed = {seed}"),
        ("真实最优（编码）", "x1 光照 = 0.7，x2 氨氮 = 0.9，x3 无机碳 = 0.6（数据构造目标）"),
        ("拟合最优（编码）", "x1 = %.3f，x2 = %.3f，x3 = %.3f（由数据分析还原）" % tuple(r["xs"])),
        ("拟合最优（实际）", "光照 %.0f μmol/m²/s，氨氮 %.0f mg/L，无机碳 %.0f mg/L"
         % (to_actual(0, r["xs"][0]), to_actual(1, r["xs"][1]), to_actual(2, r["xs"][2]))),
        ("预测最大 μmax", "%.4f d⁻¹（95%% 预测区间 %.4f ~ %.4f）"
         % (r["yopt"], r["yopt"] - r["pi_half"], r["yopt"] + r["pi_half"])),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    ws.cell(row=1, column=1).value = rows[0][0]
    ws.cell(row=1, column=1).font = Font(name="微软雅黑", bold=True, size=13, color="1F4E79")
    for i, (k, v) in enumerate(rows[1:], start=3):
        ws.cell(row=i, column=1).value = k
        ws.cell(row=i, column=1).font = Font(name="微软雅黑", bold=True, size=10.5)
        c = ws.cell(row=i, column=2); c.value = v
        c.font = BODY_FONT; c.alignment = Alignment(wrap_text=True, vertical="center")

    # --- Sheet 2: 因素与水平 ---
    ws = wb.create_sheet("因素与水平")
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12; ws.column_dimensions["F"].width = 34
    hdr = ["因子", "英文", "低水平 (-1)", "中心 (0)", "高水平 (+1)", "编码公式 x = (Z−Z₀)/ΔZ"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=1, column=j).value = h
    style_header(ws, 1, len(hdr))
    for i, f in enumerate(FACTORS):
        row = [f["name"] + f"（{f['unit']}）", f["en"], f["lo"], int(CENTER[i]), f["hi"],
               f"x = (Z − {CENTER[i]:.0f}) / {HALF[i]:.0f}"]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=2+i, column=j); c.value = v
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal="center" if j != 1 else "left",
                                    vertical="center", wrap_text=True)
    ws.cell(row=6, column=1).value = ("注：Z₀ =（低+高）/2 为中心值；ΔZ =（高−低）/2 为半幅。"
                                      "编码后 −1/0/+1 对应低/中/高水平。")
    ws.cell(row=6, column=1).font = Font(name="微软雅黑", size=9.5, color="64748B")

    # --- Sheet 3: 实验设计与结果 ---
    ws = wb.create_sheet("实验设计与结果")
    widths = [8, 8, 10, 9, 9, 9, 14, 12, 13, 13]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    hdr = ["标准序", "运行序", "点类型", "x₁ 光照\n(编码)", "x₂ 氨氮\n(编码)",
           "x₃ 无机碳\n(编码)", "光照强度\n(μmol/m²/s)", "氨氮浓度\n(mg/L)",
           "无机碳浓度\n(mg/L)", "μmax\n(d⁻¹)"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=1, column=j).value = h
    style_header(ws, 1, len(hdr))
    ws.row_dimensions[1].height = 32
    for i in range(20):
        row = 2 + i
        vals = [i + 1, int(run_no[i]), TYPES[i],
                PTS[i][0], PTS[i][1], PTS[i][2],
                to_actual(0, PTS[i][0]), to_actual(1, PTS[i][1]), to_actual(2, PTS[i][2]),
                float(Y_DATA[i])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=j)
            c.value = int(v) if isinstance(v, (int, np.integer)) else (
                float(v) if isinstance(v, (float, np.floating)) else v)
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
            if TYPES[i] in TYPE_FILL:
                c.fill = TYPE_FILL[TYPES[i]]
        ws.cell(row=row, column=10).number_format = "0.0000"
    ws.freeze_panes = "A2"

    # --- Sheet 4: 回归与方差分析 ---
    ws = wb.create_sheet("回归与方差分析")
    for j, w in enumerate([24, 12, 12, 10, 12, 10], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    r0 = 1
    ws.cell(row=r0, column=1).value = "一、二次模型回归系数（编码变量）"
    ws.cell(row=r0, column=1).font = Font(name="微软雅黑", bold=True, size=11, color="1F4E79")
    hdr = ["项", "系数", "标准误", "t 值", "p 值", "显著性"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=r0+1, column=j).value = h
    style_header(ws, r0+1, len(hdr))
    def sig_mark(p):
        return "***" if p < 0.001 else ("**" if p < 0.01 else ("*" if p < 0.05 else "ns"))
    for i in range(10):
        row = r0 + 2 + i
        vals = [TERM_NAMES[i], r["beta"][i], r["se"][i], r["t"][i], r["p"][i],
                sig_mark(r["p"][i])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=j); c.value = float(v) if j in (2, 3, 4, 5) else v
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2).number_format = "0.00000"
        ws.cell(row=row, column=3).number_format = "0.00000"
        ws.cell(row=row, column=4).number_format = "0.00"
        ws.cell(row=row, column=5).number_format = "0.000E+00"
    r0 = r0 + 13
    ws.cell(row=r0, column=1).value = "二、方差分析（ANOVA）"
    ws.cell(row=r0, column=1).font = Font(name="微软雅黑", bold=True, size=11, color="1F4E79")
    hdr = ["方差来源", "平方和 SS", "自由度 df", "均方 MS", "F 值", "p 值"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=r0+1, column=j).value = h
    style_header(ws, r0+1, len(hdr))
    anova = [
        ("回归模型", r["RegSS"], 9, r["RegSS"]/9, r["F_reg"], r["p_reg"]),
        ("残差", r["RSS"], r["dfe"], r["RSS"]/r["dfe"], None, None),
        ("  ├ 失拟 Lack of fit", r["LF"], r["dflf"], r["LF"]/r["dflf"], r["F_lof"], r["p_lof"]),
        ("  └ 纯误差 Pure error", r["PE"], r["dfpe"], r["PE"]/r["dfpe"], None, None),
        ("总计（校正）", r["TSS"], 19, None, None, None),
    ]
    for i, row_vals in enumerate(anova):
        row = r0 + 2 + i
        for j, v in enumerate(row_vals, 1):
            c = ws.cell(row=row, column=j); c.value = v
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2).number_format = "0.00000"
        ws.cell(row=row, column=4).number_format = "0.000000"
        ws.cell(row=row, column=5).number_format = "0.00"
        ws.cell(row=row, column=6).number_format = "0.000E+00"
    r0 = r0 + 8
    info = [
        ("三、拟合优度", ""),
        ("R²", "%.4f" % r["R2"]),
        ("调整 R²", "%.4f" % r["adjR2"]),
        ("残差标准差 s", "%.5f d⁻¹" % r["s"]),
        ("四、最优条件（驻点）", ""),
        ("最优编码值", "x₁ = %.3f，x₂ = %.3f，x₃ = %.3f" % tuple(r["xs"])),
        ("最优实际值", "光照 %.0f μmol/m²/s，氨氮 %.0f mg/L，无机碳 %.0f mg/L"
         % (to_actual(0, r["xs"][0]), to_actual(1, r["xs"][1]), to_actual(2, r["xs"][2]))),
        ("预测最大 μmax", "%.4f d⁻¹" % r["yopt"]),
        ("95% 置信区间（均值）", "%.4f ~ %.4f" % (r["yopt"]-r["ci_half"], r["yopt"]+r["ci_half"])),
        ("95% 预测区间（新观测）", "%.4f ~ %.4f" % (r["yopt"]-r["pi_half"], r["yopt"]+r["pi_half"])),
        ("显著性标记", "*** p<0.001，** p<0.01，* p<0.05，ns 不显著"),
    ]
    for i, (k, v) in enumerate(info):
        row = r0 + i
        ws.cell(row=row, column=1).value = k
        bold = k.startswith(("三", "四"))
        ws.cell(row=row, column=1).font = Font(name="微软雅黑", bold=True, size=10.5,
                                               color="1F4E79" if bold else "000000")
        c = ws.cell(row=row, column=2); c.value = v
        c.font = BODY_FONT
        c.alignment = Alignment(horizontal="left", vertical="center")
    p = os.path.join(BASE, "CCF_小球藻_20组模拟数据.xlsx")
    try:
        wb.save(p)
    except PermissionError:
        # 原文件被 Excel/WPS 占用时，改存 μmax 版文件名（因变量已由 Y 改为 μmax）
        p = os.path.join(BASE, "CCF_小球藻_20组模拟数据_μmax.xlsx")
        wb.save(p)
    return p

# ----------------------------------------------------------------------------
# 8. HTML 讲解报告
# ----------------------------------------------------------------------------
CSS = """
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:"Microsoft YaHei","PingFang SC",sans-serif;background:#f1f5f9;
     color:#1e293b;line-height:1.75;font-size:15px}
.wrap{max-width:960px;margin:0 auto;padding:28px 20px 60px}
header.hero{background:linear-gradient(135deg,#0f2b46 0%,#1e4e79 60%,#2d6a9f 100%);
     color:#fff;border-radius:14px;padding:30px 34px;margin-bottom:22px}
header.hero h1{font-size:25px;letter-spacing:.5px;margin-bottom:6px}
header.hero p{color:#cbd5e1;font-size:13.5px}
.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:18px 0 6px}
.card{background:rgba(255,255,255,.10);border:1px solid rgba(255,255,255,.22);
     border-radius:10px;padding:12px 14px}
.card .k{font-size:12px;color:#cbd5e1;margin-bottom:4px}
.card .v{font-size:16.5px;font-weight:700;color:#fff}
.card .s{font-size:11.5px;color:#e2e8f0;margin-top:2px}
section{background:#fff;border-radius:14px;padding:26px 30px;margin-bottom:20px;
     border:1px solid #e2e8f0}
h2{font-size:19px;color:#0f2b46;border-left:5px solid #2d6a9f;padding-left:12px;
   margin-bottom:16px}
h3{font-size:16px;color:#1e4e79;margin:18px 0 10px}
p{margin:9px 0}
li{margin:5px 0 5px 22px}
.eq{background:#f8fafc;border:1px solid #e2e8f0;border-radius:10px;padding:14px 18px;
    font-family:Cambria Math,Cambria,"Microsoft YaHei",serif;font-size:16px;
    text-align:center;margin:12px 0;color:#0f172a}
.eq .li{color:#1d4ed8}.eq .in{color:#b45309}.eq .qu{color:#047857}
table{border-collapse:collapse;width:100%;margin:12px 0;font-size:13.5px}
th{background:#1e4e79;color:#fff;padding:8px 10px;font-weight:600;text-align:center}
td{border:1px solid #e2e8f0;padding:6.5px 10px;text-align:center}
tr:nth-child(even) td{background:#f8fafc}
.tb note{display:block}
table caption{caption-side:top;text-align:left;font-size:12.5px;color:#64748b;padding-bottom:6px}
.badge{display:inline-block;padding:1px 8px;border-radius:99px;font-size:11.5px;font-weight:600}
.b-f{background:#fef3c7;color:#92400e}.b-a{background:#dbeafe;color:#1e40af}
.b-c{background:#d1fae5;color:#065f46}
.callout{border-radius:10px;padding:13px 16px;margin:13px 0;font-size:14px}
.co-key{background:#ecfdf5;border:1px solid #a7f3d0}
.co-warn{background:#fffbeb;border:1px solid #fde68a}
.co-info{background:#eff6ff;border:1px solid #bfdbfe}
.co-true{background:#f5f3ff;border:1px solid #ddd6fe}
.callout b.tag{display:block;margin-bottom:4px}
figure{margin:16px 0;text-align:center}
figure img{max-width:100%;border:1px solid #e2e8f0;border-radius:10px}
figcaption{font-size:12.5px;color:#64748b;margin-top:7px}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:14px}
footer{text-align:center;color:#94a3b8;font-size:12px;margin-top:8px}
.num{font-variant-numeric:tabular-nums}
.pos{color:#b45309;font-weight:600}.neg{color:#1d4ed8;font-weight:600}
.toc{background:#fff;border:1px solid #e2e8f0;border-radius:14px;padding:16px 24px;
     margin-bottom:20px;font-size:13.5px}
.toc a{color:#1e4e79;text-decoration:none;margin-right:16px}
.small{font-size:12.5px;color:#64748b}
@media print{body{background:#fff}section{border:none;padding:10px 0}}
"""

def fmt_p(p):
    return "&lt; 0.001" if p < 0.001 else f"{p:.3f}"

def img_tag(path, cap, w="100%"):
    b64 = base64.b64encode(open(path, "rb").read()).decode()
    return (f'<figure><img src="data:image/png;base64,{b64}" style="width:{w}">'
            f"<figcaption>{cap}</figcaption></figure>")

def make_html(r, seed, sigma, run_no, figs):
    b = r["beta"]
    # 真实系数换算到 μmax 标度（仿射标定：截距 b + a·B0，其余 a·B）
    true_b = [B_SCL + A_SCL * B0] + [A_SCL * v for v in (*B_LIN, *B_INT, *B_QUAD)]
    eq_fit = (f"ŷ(μmax) = {b[0]:.4f} + <span class='li'>{b[1]:.4f}x₁</span> + "
              f"<span class='li'>{b[2]:.4f}x₂</span> + <span class='li'>{b[3]:.4f}x₃</span> "
              f"+ <span class='in'>{b[4]:.4f}x₁x₂</span> + "
              f"<span class='in'>{b[5]:.4f}x₁x₃</span> + "
              f"<span class='in'>{b[6]:.4f}x₂x₃</span> "
              f"− <span class='qu'>{abs(b[7]):.4f}x₁²</span> − "
              f"<span class='qu'>{abs(b[8]):.4f}x₂²</span> − "
              f"<span class='qu'>{abs(b[9]):.4f}x₃²</span>")
    eq_true = (f"μmax = {true_b[0]:.4f} + {true_b[1]:.4f}x₁ + {true_b[2]:.4f}x₂ "
               f"+ {true_b[3]:.4f}x₃ + {true_b[4]:.4f}x₁x₂ + {true_b[5]:.4f}x₁x₃ "
               f"− {abs(true_b[6]):.4f}x₂x₃ − {abs(true_b[7]):.4f}x₁² "
               f"− {abs(true_b[8]):.4f}x₂² − {abs(true_b[9]):.4f}x₃² + ε")

    # 20 组数据表
    trs = []
    for i in range(20):
        cls = {"析因点": "b-f", "轴点": "b-a", "中心点": "b-c"}[TYPES[i]]
        trs.append(
            f"<tr><td class='num'>{i+1}</td><td class='num'>{int(run_no[i])}</td>"
            f"<td><span class='badge {cls}'>{TYPES[i]}</span></td>"
            f"<td class='num'>{int(PTS[i][0]):+d}</td><td class='num'>{int(PTS[i][1]):+d}</td>"
            f"<td class='num'>{int(PTS[i][2]):+d}</td>"
            f"<td class='num'>{to_actual(0, PTS[i][0]):.0f}</td>"
            f"<td class='num'>{to_actual(1, PTS[i][1]):.0f}</td>"
            f"<td class='num'>{to_actual(2, PTS[i][2]):.0f}</td>"
            f"<td class='num'><b>{Y_DATA[i]:.4f}</b></td></tr>")
    table20 = ("<table><caption>标准序 1–8 为析因点（全部 ±1 组合）；9–14 为轴点（面心，α=1）；"
               "15–20 为中心点（0,0,0 重复 6 次）。运行序为随机化执行顺序。</caption>"
               "<tr><th>标准序</th><th>运行序</th><th>点类型</th><th>x₁ 光照<br>(编码)</th>"
               "<th>x₂ 氨氮<br>(编码)</th><th>x₃ 无机碳<br>(编码)</th>"
               "<th>光照强度<br>(μmol/m²/s)</th><th>氨氮浓度<br>(mg/L)</th>"
               "<th>无机碳浓度<br>(mg/L)</th><th>μmax<br>(d⁻¹)</th></tr>"
               + "".join(trs) + "</table>")

    # 系数表
    coef_rows = []
    for i in range(10):
        sig = "***" if r["p"][i] < 0.001 else ("**" if r["p"][i] < 0.01 else
              ("*" if r["p"][i] < 0.05 else "ns"))
        coef_rows.append(f"<tr><td>{TERM_NAMES[i]}</td><td class='num'>{b[i]:.5f}</td>"
                         f"<td class='num'>{r['se'][i]:.5f}</td>"
                         f"<td class='num'>{r['t'][i]:.2f}</td>"
                         f"<td class='num'>{fmt_p(r['p'][i])}</td>"
                         f"<td><b>{sig}</b></td></tr>")
    table_coef = ("<table><tr><th>项</th><th>系数</th><th>标准误</th><th>t 值</th>"
                  f"<th>p 值</th><th>显著性</th></tr>" + "".join(coef_rows) + "</table>")

    # 真实 vs 拟合对照（真实系数已换算到 μmax 标度）
    cmp_rows = []
    for i in range(10):
        cmp_rows.append(f"<tr><td>{TERM_NAMES[i]}</td><td class='num'>{true_b[i]:+.5f}</td>"
                        f"<td class='num'>{b[i]:+.5f}</td>"
                        f"<td class='num'>{b[i]-true_b[i]:+.5f}</td></tr>")
    table_cmp = ("<table><tr><th>项</th><th>真实系数（构造值）</th><th>拟合系数（数据还原）</th>"
                 "<th>偏差</th></tr>" + "".join(cmp_rows) + "</table>")

    # ANOVA
    def ssf(v): return f"{v:.5f}"
    def msf(v): return f"{v:.6f}"
    table_anova = f"""
<table><tr><th>方差来源</th><th>平方和 SS</th><th>自由度 df</th><th>均方 MS</th>
<th>F 值</th><th>p 值</th></tr>
<tr><td>回归模型</td><td class='num'>{ssf(r['RegSS'])}</td><td>9</td>
<td class='num'>{msf(r['RegSS']/9)}</td><td class='num'>{r['F_reg']:.1f}</td>
<td class='num'>{fmt_p(r['p_reg'])}</td></tr>
<tr><td>残差</td><td class='num'>{ssf(r['RSS'])}</td><td>{r['dfe']}</td>
<td class='num'>{msf(r['RSS']/r['dfe'])}</td><td colspan="2">s = {r['s']:.5f} d⁻¹</td></tr>
<tr><td>&nbsp;&nbsp;├ 失拟（模型偏差）</td><td class='num'>{ssf(r['LF'])}</td>
<td>{r['dflf']}</td><td class='num'>{msf(r['LF']/r['dflf'])}</td>
<td class='num'>{r['F_lof']:.2f}</td><td class='num'>{fmt_p(r['p_lof'])}</td></tr>
<tr><td>&nbsp;&nbsp;└ 纯误差（中心点重复）</td><td class='num'>{ssf(r['PE'])}</td>
<td>{r['dfpe']}</td><td class='num'>{msf(r['PE']/r['dfpe'])}</td>
<td colspan="2">中心点均值 {r['center_mean']:.4f} ± {r['center_sd']:.4f}</td></tr>
<tr><td><b>总计（校正）</b></td><td class='num'>{ssf(r['TSS'])}</td><td>19</td>
<td colspan="3">R² = <b>{r['R2']:.4f}</b>，调整 R² = <b>{r['adjR2']:.4f}</b></td></tr>
</table>"""

    # 边际效应（交互作用的数字证据，来自拟合模型）
    dL_N_lo, dL_N_hi = 2*(b[1]-b[4]), 2*(b[1]+b[4])     # 光照全程效应 @低氮/高氮
    dL_C_lo, dL_C_hi = 2*(b[1]-b[5]), 2*(b[1]+b[5])     # 光照全程效应 @低碳/高碳
    dC_N_lo, dC_N_hi = 2*(b[3]-b[6]), 2*(b[3]+b[6])     # 无机碳全程效应 @低氮/高氮
    dN_L_lo, dN_L_hi = 2*(b[2]-b[4]), 2*(b[2]+b[4])     # 氨氮全程效应 @低光/高光
    xs = r["xs"]
    act_opt = [to_actual(0, xs[0]), to_actual(1, xs[1]), to_actual(2, xs[2])]

    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>CCF 设计讲解报告 —— 小球藻 μmax 响应面模拟数据</title>
<style>{CSS}</style></head><body><div class="wrap">

<header class="hero">
<h1>中心复合面心设计（CCF）讲解报告 · μmax 响应面</h1>
<p>小球藻（Chlorella sp.）微反应器培养 · 响应：最大比生长速率 μmax · 三因子响应面模拟数据集（20 组）· 生成于
{datetime.now().strftime('%Y-%m-%d')}</p>
<div class="cards">
<div class="card"><div class="k">最优光照（编码 0.7）</div>
<div class="v">{act_opt[0]:.0f} μmol/m²/s</div><div class="s">x₁* = {xs[0]:.3f}（拟合）</div></div>
<div class="card"><div class="k">最优氨氮（编码 0.9）</div>
<div class="v">{act_opt[1]:.0f} mg/L</div><div class="s">x₂* = {xs[1]:.3f}（拟合）</div></div>
<div class="card"><div class="k">最优无机碳（编码 0.6）</div>
<div class="v">{act_opt[2]:.0f} mg/L</div><div class="s">x₃* = {xs[2]:.3f}（拟合）</div></div>
<div class="card"><div class="k">预测最大 μmax</div>
<div class="v">{r['yopt']:.4f} d⁻¹</div>
<div class="s">95% PI：{r['yopt']-r['pi_half']:.4f} ~ {r['yopt']+r['pi_half']:.4f}</div></div>
</div>
</header>

<div class="toc"><b>目录</b>：<a href="#s1">① 任务定义</a><a href="#s2">② CCF 设计原理</a>
<a href="#s3">③ 20 组模拟数据</a><a href="#s4">④ 模型拟合与方差分析</a>
<a href="#s5">⑤ 交互作用详解</a><a href="#s6">⑥ 最优值求解与含义</a>
<a href="#s7">⑦ 实践要点小结</a></div>

<section id="s1"><h2>① 研究任务定义</h2>
<p>在微藻微反应器中培养小球藻，考察三个操作因子对<b>最大比生长速率 μmax（d⁻¹）</b>的影响
（μmax 由批式培养的生长曲线拟合得到，衡量细胞增殖的“速度上限”），并寻找使 μmax 最大的培养条件：</p>
<table>
<tr><th>因子</th><th>符号</th><th>范围</th><th>低 (−1)</th><th>中心 (0)</th><th>高 (+1)</th><th>半幅 ΔZ</th></tr>
<tr><td>光照强度 Light intensity</td><td>x₁</td><td>30–210 μmol/m²/s</td><td>30</td>
<td>120</td><td>210</td><td>90</td></tr>
<tr><td>氨氮浓度 NH₄⁺-N</td><td>x₂</td><td>20–300 mg/L</td><td>20</td><td>160</td>
<td>300</td><td>140</td></tr>
<tr><td>无机碳浓度 IC</td><td>x₃</td><td>75–565 mg/L</td><td>75</td><td>320</td>
<td>565</td><td>245</td></tr>
</table>
<p>要求：（1）采用 <b>CCF 设计</b>共 20 组；（2）数据的<b>最优编码水平</b>为 x₁=0.7、x₂=0.9、
x₃=0.6；（3）三对因子间均存在<b>显著且符合培养机理</b>的交互作用；（4）20 组数据的 μmax
<b>最小值 = 0.0571 d⁻¹、最大值 = 0.2072 d⁻¹</b>（精确锁定）。</p>
<div class="callout co-true"><b class="tag">数据构造说明（教学透明度）</b>
本数据集由“真实二次模型 + 正态测量误差（σ ≈ {A_SCL*sigma:.4f} d⁻¹）”模拟生成：先在基准标度上
搜索满足全部统计要求的随机实现（驻点贴近 0.7/0.9/0.6、交互与二次项显著、失拟不显著），
再经<b>仿射标定 y′ = a·y + b</b>（a = {A_SCL:.6f}，b = {B_SCL:.6f}）把 20 个观测精确锁定到
[0.0571, 0.2072] d⁻¹。<b>仿射标定不改变驻点位置，也不改变任何 t 值、p 值、R² 与失拟检验结论</b>
（系数与噪声等比例缩放，检验统计量是尺度不变的）——因此最优条件与交互作用显著性原样保留。
随后<b>只用数据</b>走完整个 RSM 分析流程；构造目标与拟合结果对照见第 ④ 节。</div>
</section>

<section id="s2"><h2>② CCF 设计原理：为什么是这 20 组</h2>
<h3>2.1 从 CCD 家族说起</h3>
<p>响应面法（RSM）用<b>二次多项式</b>描述因子与响应的关系。3 个因子时模型含
1 个截距 + 3 个线性项 + 3 个交互项 + 3 个二次项 = <b>10 个参数</b>，因此至少需要 10 个
不同试验条件，且每个因子至少 3 个水平（否则估不出二次项）。</p>
<p><b>中心复合设计（CCD）</b>是 RSM 的标准方案，由三类点组成：</p>
<table>
<tr><th>点的类型</th><th>个数（k=3）</th><th>位置</th><th>作用</th></tr>
<tr><td>析因点</td><td>2³ = 8</td><td>立方体 8 个角（±1 组合）</td>
<td>估计线性项与交互项</td></tr>
<tr><td>轴点（星点）</td><td>2k = 6</td><td>距中心 ±α（其余为 0）</td>
<td>估计二次项、撑开设计空间</td></tr>
<tr><td>中心点</td><td>6（本设计）</td><td>(0,0,0)</td>
<td>估计纯误差、失拟检验、稳定性监测</td></tr>
</table>
<p><b>α 的取值决定 CCD 的类型</b>：α = (2^k)^¼ ≈ 1.68 为“可旋转 CCD”（各方向预测精度一致）；
<b>α = 1 时轴点恰好落在立方体六个面的中心，即 CCF（面心设计）</b>。</p>

<h3>2.2 为什么本实验必须用 CCF</h3>
<div class="callout co-warn"><b class="tag">关键：可旋转 CCD 在本题的因子范围下物理不可行</b>
可旋转 CCD 的星臂要伸到 ±1.68 个半幅之外，代入本题三个因子：
光照 120 ± 1.68×90 = <b>−31 ~ 271 μmol/m²/s</b>（负光照不存在）；<br>
氨氮 160 ± 1.68×140 = <b>−76 ~ 396 mg/L</b>（负浓度不存在）；<br>
无机碳 320 ± 1.68×245 = <b>−92 ~ 732 mg/L</b>（负浓度不存在）。<br>
三个因子的下限全部撞上物理边界（0），因此只能取 α=1，把轴点放在面上。
<b>代价</b>：设计不旋转、模型不能外推到立方体之外；<b>收益</b>：每因子只需 3 个真实水平、
全部条件可实现、设备操作最简单——这正是 CCF 在生物培养类实验中流行的原因。</div>
{img_tag(figs['fig1'], '图 1 CCF 设计空间。橙色圆点 = 8 个析因点（角），蓝色方块 = 6 个轴点（面心，α=1），绿色星 = 中心点（重复 6 次）。三个因子的编码 −1/0/+1 对应各自的实际水平。')}
<h3>2.3 编码：把实际量纲变成可比较的数</h3>
<div class="eq">x = (Z − Z₀) / ΔZ ，其中 Z₀ =（Z<sub>低</sub>+Z<sub>高</sub>）/2（中心值），
ΔZ =（Z<sub>高</sub>−Z<sub>低</sub>）/2（半幅）</div>
<p>例如光照 183 μmol/m²/s → x₁ = (183−120)/90 = <b>0.7</b>。编码后所有因子同尺度（−1~+1），
系数大小可直接比较，也避免了量纲悬殊带来的数值问题。</p>
<h3>2.4 6 个中心点不是浪费</h3>
<p>① 重复中心点给出<b>纯误差</b>估计（本例 6 点 → 5 个自由度），这是“失拟检验”的基准；
② 检验<b>弯曲效应</b>——若中心点均值明显偏离角点/轴点的线性外推，说明必须用二次模型；
③ 随机穿插在运行序列中，可监测系统漂移；④ 提高最感兴趣区域（中心区）的预测精度。</p>
</section>

<section id="s3"><h2>③ 20 组模拟数据</h2>
<p>真实模型（构造用，编码变量，已换算到 μmax 标度）：</p>
<div class="eq">{eq_true}</div>
{table20}
<p class="small">μmax 保留四位小数（模拟测量精度 0.0001 d⁻¹），最小值 0.0571、最大值 0.2072 被
精确锁定；最小值出现在“强光 + 低氮 + 低碳”组合而非全低组合——这是低氮下<b>光抑制</b>的直接体现，
详见第 ⑤ 节。</p>
</section>

<section id="s4"><h2>④ 模型拟合与统计检验</h2>
<p>对 20 组数据做全二次模型最小二乘拟合（编码变量），得到：</p>
<div class="eq">{eq_fit}</div>
<p><span class="li">蓝色 = 线性项</span>，<span class="in">橙色 = 交互项</span>，
<span class="qu">绿色 = 二次项</span>。读法示例：交互项系数
x₁x₂ = {b[4]:.4f} 表示<b>氨氮每升高 1 个编码单位，光照的斜率（边际效应）就增加
{b[4]:.4f} d⁻¹</b>——这就是“交互”的数学含义。</p>
{table_coef}
{table_anova}
<div class="callout co-key"><b class="tag">如何读这张 ANOVA 表</b>
① 回归模型整体极显著（p {fmt_p(r['p_reg'])}）→ 模型有效；
② <b>失拟 p = {r['p_lof']:.3f} > 0.05</b> → 二次模型已充分刻画系统，无未建模的系统性偏差；
③ 三个<b>交互项与三个二次项全部 p&lt;0.01</b> → 交互作用与非线性均真实存在；
④ R² = {r['R2']:.3f}、调整 R² = {r['adjR2']:.3f} → 模型解释了
{100*r['R2']:.1f}% 的数据波动；残差标准差 s = {r['s']:.5f} d⁻¹，与标定后构造噪声
σ ≈ {A_SCL*sigma:.4f} d⁻¹ 一致。</div>
{img_tag(figs['fig5'], '图 2 效应 Pareto 图。三对交互作用全部显著，且强度排序符合培养机理：氮的主效应最强（氮是限制性营养盐）。', '88%')}
{img_tag(figs['fig6'], '图 3 观测值 vs 预测值。数据点均匀紧贴 45° 线，无系统性弯曲或离群——模型无偏、方差齐性良好的直观证据。', '58%')}
<h3>拟合 vs 真实：分析能否还原构造模型？</h3>
{table_cmp}
<p class="small">所有系数偏差均在 ±{max(abs(b[i]-true_b[i]) for i in range(10)):.4f} 之内
（与 s/√n 的理论标准误一致），驻点还原为 ({xs[0]:.3f}, {xs[1]:.3f}, {xs[2]:.3f})，
与构造目标 (0.7, 0.9, 0.6) 的最大偏差 {np.abs(r['xs']-TARGET).max():.3f} 个编码单位。</p>
</section>

<section id="s5"><h2>⑤ 交互作用详解：折线不平行之处，就是因子在“合作或打架”</h2>
<p><b>判读规则</b>：交互图横轴为因子 A，不同折线为因子 B 的水平。若折线<b>平行</b> →
B 不改变 A 的效应（无交互）；<b>不平行</b> → A 的效应大小取决于 B 的水平（有交互）。
张口越大、越不平行，交互越强。</p>
{img_tag(figs['fig2'], '图 4 三对交互作用图（基于拟合模型，第三因子固定在中心水平）。三组折线明显不平行，交互作用方向各不相同：协同、协同、拮抗。')}

<h3>5.1 光照 × 氨氮 —— <span class="pos">正交互（协同）</span>：
没有氮，“接不住”光</h3>
<p>机理：氮是叶绿素分子与光合系统蛋白的组成元素。<b>低氮</b>下细胞无法合成足够的光合机构，
此时强光非但无法被利用，还会产生活性氧引起<b>光抑制/光漂白</b>；<b>高氮</b>下细胞“接得住”
光，光照的增产效应才能完全释放。</p>
<p>数据证据（拟合模型，无机碳固定 320 mg/L）：光照从 30 提高到 210 μmol/m²/s，
低氮（20 mg/L）时 μmax 变化 ≈ <span class="neg">{dL_N_lo:+.4f} d⁻¹</span>（几乎为零、甚至微降——
光抑制）；高氮（300 mg/L）时 ≈ <span class="pos">+{dL_N_hi:.4f} d⁻¹</span>。
对应地，全设计<b>最低</b> μmax（0.0571 d⁻¹）出现在（+1, −1, −1）即“强光 + 低氮 + 低碳”，
而不是全低组合——教科书式的低氮光抑制。</p>

<h3>5.2 光照 × 无机碳 —— <span class="pos">正交互（协同）</span>：
能量与底物必须配套</h3>
<p>机理：光反应提供能量（ATP/NADPH），无机碳是卡尔文循环的底物。只提高其中一个，
反应立刻撞上另一个的瓶颈（Liebig 最小因子定律）；两者同步提升，光合速率才能台阶式上升。</p>
<p>数据证据（氨氮固定 160 mg/L）：光照 30→210 的效应在低碳（75 mg/L）时仅
≈ <span class="neg">{dL_C_lo:+.4f} d⁻¹</span>，在高碳（565 mg/L）时达
<span class="pos">+{dL_C_hi:.4f} d⁻¹</span>——碳充足时补光才有效，反之亦然。</p>

<h3>5.3 氨氮 × 无机碳 —— <span class="neg">负交互（拮抗）</span>：
高氮时加碳要克制</h3>
<p>机理：(i) <b>收益递减</b>——高氮背景已接近碳利用饱和，继续增碳边际收益下滑；
(ii) <b>游离氨毒性</b>——以碳酸氢盐补碳会推高培养液 pH，使 NH₄⁺ 向可自由穿透细胞膜的
游离 NH₃ 转变，解耦联光合磷酸化。“高氮 + 高碳”组合互相拖累，这正是系数为负的原因。</p>
<p>数据证据（光照固定 120 μmol/m²/s）：无机碳 75→565 的效应在低氮时
≈ <span class="pos">+{dC_N_lo:.4f} d⁻¹</span>，在高氮时缩水到
≈ <span class="neg">+{dC_N_hi:.4f} d⁻¹</span>。<b>工艺启示</b>：最优碳水平在高氮时应主动调低
（最优 x₃ = {xs[2]:.2f} &lt; 1 正体现了这一点），负交互直接改写操作策略。</p>
<div class="callout co-info"><b class="tag">三对交互的生物学逻辑链</b>
光照×氨氮（+）：氮是光合机构的“建材”，建材不足时光照是毒药；<br>
光照×无机碳（+）：光是能量、碳是底物，二者是光合作用的“配套件”；<br>
氨氮×无机碳（−）：高铵 + 高碱度 → 游离氨胁迫，且高氮下碳收益饱和。</div>
</section>

<section id="s6"><h2>⑥ 最优条件：0.7 / 0.9 / 0.6 到底是什么意思</h2>
<h3>6.1 驻点求解</h3>
<p>对拟合模型令三个偏导数为零，解 3×3 线性方程组：</p>
<div class="eq">∂μmax/∂x₁ = {b[1]:.4f} + {2*b[7]:.4f}x₁ + {b[4]:.4f}x₂ + {b[5]:.4f}x₃ = 0<br>
∂μmax/∂x₂ = {b[2]:.4f} + {b[4]:.4f}x₁ + {2*b[8]:.4f}x₂ + {b[6]:.4f}x₃ = 0<br>
∂μmax/∂x₃ = {b[3]:.4f} + {b[5]:.4f}x₁ + {b[6]:.4f}x₂ + {2*b[9]:.4f}x₃ = 0</div>
<p>解得驻点 <b>x₁* = {xs[0]:.3f}，x₂* = {xs[1]:.3f}，x₃* = {xs[2]:.3f}</b>。
Hessian 矩阵三个特征值均为负（{', '.join(f'{v:.4f}' for v in sorted(r['eig'], reverse=True))}）
→ 驻点是<b>极大值</b>（山顶），而非鞍点。</p>
<h3>6.2 编码 → 实际：把 0.7 翻译回工程语言</h3>
<table>
<tr><th>因子</th><th>最优编码值</th><th>换算</th><th>最优实际值</th></tr>
<tr><td>光照强度</td><td><b>{xs[0]:.2f}</b></td><td>120 + {xs[0]:.2f}×90</td>
<td><b>{act_opt[0]:.0f} μmol/m²/s</b></td></tr>
<tr><td>氨氮浓度</td><td><b>{xs[1]:.2f}</b></td><td>160 + {xs[1]:.2f}×140</td>
<td><b>{act_opt[1]:.0f} mg/L</b></td></tr>
<tr><td>无机碳浓度</td><td><b>{xs[2]:.2f}</b></td><td>320 + {xs[2]:.2f}×245</td>
<td><b>{act_opt[2]:.0f} mg/L</b></td></tr>
</table>
<div class="callout co-key"><b class="tag">“最优 = 0.7”的三层含义</b>
① <b>几何</b>：0.7 = 从中心点(0)朝高水平(+1)方向走了 70% 的路程；
② <b>实验设计</b>：三个最优编码值都严格落在 (−1, +1) <b>内部</b> → 实验范围把真实的
“山顶”完整罩住了，范围选得好；若某因子的最优编码贴近 +1 甚至超出 +1，说明真最优在
范围之外，需要扩范围补做实验；
③ <b>模型边界</b>：CCF 的 α=1 决定了模型只在立方体内可靠——本例驻点在立方体内部
（0.7/0.9/0.6 均未触及 ±1），因此最优点预测可信、无需外推。</div>
{img_tag(figs['fig3'], '图 5 等高线图。每张图把第三个因子固定在其最优编码值，★ 为最优点。可见山顶被设计空间完整覆盖；等高线呈同心“椭圆”状（受交互项影响略倾斜），靠近边缘处响应下降加快。')}
{img_tag(figs['fig4'], '图 6 响应面三维图（无机碳固定于最优水平）。曲面在 (0.7, 0.9) 附近达到峰值 μmax ≈ %.4f d⁻¹。' % r['yopt'], '76%')}
<h3>6.3 最优点的预测与验证</h3>
<table>
<tr><th>项目</th><th>数值</th></tr>
<tr><td>预测最大 μmax（点估计）</td><td><b>{r['yopt']:.4f} d⁻¹</b></td></tr>
<tr><td>95% 置信区间（最优条件下的均值）</td>
<td>{r['yopt']-r['ci_half']:.4f} ~ {r['yopt']+r['ci_half']:.4f} d⁻¹</td></tr>
<tr><td>95% 预测区间（下次单次实验）</td>
<td>{r['yopt']-r['pi_half']:.4f} ~ {r['yopt']+r['pi_half']:.4f} d⁻¹</td></tr>
</table>
<p class="small">预测峰值 {r['yopt']:.4f} d⁻¹ 略高于 20 组观测的最大值 0.2072 d⁻¹——峰顶在
设计空间内部，没有任何实验点恰好踩在山顶上，这正是需要用模型解驻点方程（而非直接挑数据最大值）
的原因。</p>
<div class="callout co-warn"><b class="tag">确认实验（response validation）——RSM 的收尾动作</b>
在最优条件（光照 {act_opt[0]:.0f} μmol/m²/s、氨氮 {act_opt[1]:.0f} mg/L、
无机碳 {act_opt[2]:.0f} mg/L）下做 <b>3 次独立重复</b>。若实测均值落入 95% 预测区间
（{r['yopt']-r['pi_half']:.4f} ~ {r['yopt']+r['pi_half']:.4f} d⁻¹），模型即验证通过，
该条件可作为推荐工艺；若明显偏离，需检查是否存在未控制变量（温度、pH 漂移、接种批次等）
并考虑在最优区域做补充设计（如再补一组小范围 CCF）。</div>
</section>

<section id="s7"><h2>⑦ 实践要点小结</h2>
<p><b>CCF 全流程清单（可复制到任何 3 因子优化问题）：</b></p>
<li>① 用文献 + 预实验确定因子范围，中心尽量放在“猜测的最优”附近；</li>
<li>② 检查范围能否容纳可旋转 CCD 的 ±1.68 半幅星臂（本题不能，负水平物理不存在）→ 选 CCF；</li>
<li>③ 生成 8 析因 + 6 轴点 + 5~6 中心点，<b>随机化运行顺序</b>（必要时按天分区块）；</li>
<li>④ 实验测定响应（本例为模拟数据）；</li>
<li>⑤ 拟合全二次模型 → 看 ANOVA：整体显著性、失拟 p&gt;0.05、R²/调整 R²；</li>
<li>⑥ 逐项检验：不显著的高阶项可酌情合并进误差（本例全部显著，保留完整模型）；</li>
<li>⑦ 解驻点方程 → 判断极大/鞍点 → 编码换算回实际值；</li>
<li>⑧ 画等高线/响应面确认山顶在设计空间内部；</li>
<li>⑨ 在预测最优点做确认实验（3 重复），落入预测区间才算闭环。</li>
<p class="small">注意事项：CCF 不能外推到立方体外；中心点分散随机化以监测漂移；
若因子间存在强非线性（如抑制剂的陡峭毒性峰），二次模型可能失拟，需换更高阶设计
（如 Box-Behnken 之外的专用设计）或分段建模。</p>
</section>

<footer>CCF 模拟数据集（μmax 标度）· seed={seed}，σ ≈ {A_SCL*sigma:.4f} d⁻¹ ·
仿射标定 a = {A_SCL:.6f}、b = {B_SCL:.6f} · 生成于
{datetime.now().strftime('%Y-%m-%d %H:%M')} · 分析脚本：ccf_simulate.py</footer>
</div></body></html>"""
    p = os.path.join(BASE, "CCF_小球藻_讲解报告.html")
    with open(p, "w", encoding="utf-8") as f:
        f.write(html)
    return p

# ----------------------------------------------------------------------------
# 9. 主流程
# ----------------------------------------------------------------------------
PTS, TYPES = build_design()

def main():
    print("=" * 70)
    print("CCF 模拟数据生成 + RSM 分析（响应：μmax，范围锁定 [0.0571, 0.2072] d⁻¹）")
    xs_true = verify_true_stationary()
    print(f"[1] 真实模型驻点校验通过: {xs_true} （Hessian 负定 → 极大值）")
    print(f"[2] 真实模型最优实际值: 光照 {to_actual(0, .7):.0f} μmol/m²/s, "
          f"氨氮 {to_actual(1, .9):.0f} mg/L, 无机碳 {to_actual(2, .6):.0f} mg/L")
    print("[3] 种子搜索（基准标度）...")
    global Y_DATA, A_SCL, B_SCL
    seed, sigma, y_base = seed_search()
    # 仿射标定：把基准标度数据精确锁定到 μmax ∈ [0.0571, 0.2072]
    lo, hi = float(y_base.min()), float(y_base.max())
    A_SCL = (TARGET_MAX - TARGET_MIN) / (hi - lo)
    B_SCL = TARGET_MIN - A_SCL * lo
    Y_DATA = np.round(A_SCL * y_base + B_SCL, 4)
    assert abs(float(Y_DATA.min()) - TARGET_MIN) < 1e-9, f"min={Y_DATA.min()}"
    assert abs(float(Y_DATA.max()) - TARGET_MAX) < 1e-9, f"max={Y_DATA.max()}"
    print(f"[4] 仿射标定: a={A_SCL:.6f}, b={B_SCL:.6f}; "
          f"μmax 范围 [{Y_DATA.min():.4f}, {Y_DATA.max():.4f}] d⁻¹ 已精确锁定")
    print("    μmax(20):", ", ".join(f"{v:.4f}" for v in Y_DATA))
    # 运行顺序（独立可复现）
    order = np.random.default_rng(seed + 99991).permutation(20)
    run_no = np.empty(20, dtype=int)
    for pos, idx in enumerate(order):
        run_no[idx] = pos + 1
    print("    中心点均值 %.4f, SD %.4f (d⁻¹)" %
          (Y_DATA[14:20].mean(), Y_DATA[14:20].std(ddof=1)))

    r = analyze(PTS, Y_DATA)
    # 标定后统计量复核（仿射不变性 + 4 位舍入的鲁棒性）
    assert max(r["p"][4], r["p"][5], r["p"][6]) < 0.01, "交互项显著性丢失"
    assert max(r["p"][7], r["p"][8], r["p"][9]) < 0.01, "二次项显著性丢失"
    assert r["p_lof"] > 0.10, "失拟变得显著"
    assert r["eig"].max() < 0, "驻点不再是极大值"
    assert np.abs(r["xs"] - TARGET).max() <= 0.08, "驻点偏离目标"
    print(f"[5] 标定后复核通过: R2={r['R2']:.4f}, adjR2={r['adjR2']:.4f}, "
          f"s={r['s']:.5f} d⁻¹")
    print(f"    交互项 p: x1x2={r['p'][4]:.2e}, x1x3={r['p'][5]:.2e}, x2x3={r['p'][6]:.2e}")
    print(f"    失拟 p={r['p_lof']:.3f}")
    print(f"    拟合驻点: {np.round(r['xs'], 4)}（目标 {TARGET}）")
    print(f"    预测最优 μmax: {r['yopt']:.4f} d⁻¹")

    print("[6] 绘图 ...")
    figs = {}
    figs["fig1"] = fig1_design()
    figs["fig2"] = fig2_interaction(r["beta"])
    figs["fig3"] = fig3_contour(r["beta"], r["xs"])
    figs["fig4"] = fig4_surface(r["beta"], r["xs"])
    figs["fig5"] = fig5_pareto(r)
    figs["fig6"] = fig6_obspred(r["beta"], r)
    for k, v in figs.items():
        print(f"    {k}: {os.path.basename(v)}")

    print("[7] 交互作用 SVG ...")
    svg_p = svg_interaction(r["beta"])
    print("    " + os.path.basename(svg_p))

    print("[8] Excel ...")
    xlsx_p = make_xlsx(r, seed, sigma, run_no)
    print("    " + os.path.basename(xlsx_p))

    print("[9] HTML 报告 ...")
    html_p = make_html(r, seed, sigma, run_no, figs)
    print("    " + os.path.basename(html_p))

    summary = dict(
        seed=int(seed), sigma=float(sigma),
        response=dict(name=RESP_NAME, unit=RESP_UNIT),
        scale=dict(a=float(A_SCL), b=float(B_SCL), target_min=TARGET_MIN,
                   target_max=TARGET_MAX, sigma_effective=float(A_SCL*sigma)),
        design=[dict(std=i+1, run=int(run_no[i]), type=TYPES[i],
                     x1=float(PTS[i][0]), x2=float(PTS[i][1]), x3=float(PTS[i][2]),
                     light=to_actual(0, PTS[i][0]), nh4=to_actual(1, PTS[i][1]),
                     ic=to_actual(2, PTS[i][2]), y=float(Y_DATA[i])) for i in range(20)],
        beta=[float(v) for v in r["beta"]],
        se=[float(v) for v in r["se"]],
        t=[float(v) for v in r["t"]],
        p=[float(v) for v in r["p"]],
        R2=float(r["R2"]), adjR2=float(r["adjR2"]), s=float(r["s"]),
        p_lof=float(r["p_lof"]), F_lof=float(r["F_lof"]),
        xs=[float(v) for v in r["xs"]], eig=[float(v) for v in r["eig"]],
        yopt=float(r["yopt"]), pi=[float(r["yopt"]-r["pi_half"]), float(r["yopt"]+r["pi_half"])],
        ci=[float(r["yopt"]-r["ci_half"]), float(r["yopt"]+r["ci_half"])],
        files=dict(xlsx=xlsx_p, html=html_p, svg=svg_p,
                   figures={k: v for k, v in figs.items()}),
    )
    with open(os.path.join(BASE, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=1)
    print("[10] DONE")

if __name__ == "__main__":
    main()

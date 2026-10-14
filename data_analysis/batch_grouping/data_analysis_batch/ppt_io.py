from __future__ import annotations

import io
import json
import math
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from .grouping import metrics, sort_for_output
from .models import ChamberResult


def _target_image(worksheet):
    for image in worksheet._images:
        anchor = image.anchor
        if hasattr(anchor, "_from") and anchor._from.col == 7 and anchor._from.row == 1:
            return image
    if worksheet._images:
        return worksheet._images[0]
    raise RuntimeError(f"No embedded growth curve image found in worksheet {worksheet.title}")


def _quality(image_bytes: bytes) -> tuple[int, int, float, float]:
    with Image.open(io.BytesIO(image_bytes)) as image:
        width, height = image.size
        raw_dpi = image.info.get("dpi", (0, 0))
        if isinstance(raw_dpi, (int, float)):
            dpi_x = dpi_y = float(raw_dpi)
        else:
            dpi_x, dpi_y = float(raw_dpi[0]), float(raw_dpi[1])
    return width, height, dpi_x, dpi_y


def extract_growth_charts(
    base_dir: Path,
    work_dir: Path,
    groups: Sequence[Sequence[ChamberResult]],
) -> tuple[Path, dict]:
    chart_root = work_dir / "ppt_charts"
    chart_root.mkdir(parents=True, exist_ok=True)
    for path in chart_root.glob("group_*"):
        if path.is_dir():
            shutil.rmtree(path)

    manifest: dict = {"groups": []}
    tasks: dict[str, list[tuple[ChamberResult, Path, dict]]] = defaultdict(list)
    for group_index, group in enumerate(groups, start=1):
        ordered = sort_for_output(group)
        result = metrics(group)
        group_dir = chart_root / f"group_{group_index}"
        group_dir.mkdir(parents=True, exist_ok=True)
        group_manifest = {
            "index": group_index,
            "size": len(ordered),
            "source_workbook": f"selected__groups_{group_index}_{len(ordered)}.xlsx",
            "selected_sheet": f"selected_{len(ordered)}",
            "umax_mean_h-1": result.mean,
            "umax_cv_percent": result.cv * 100,
            "charts": [],
        }
        for order, row in enumerate(ordered, start=1):
            output_path = group_dir / f"{order:02d}_{row.chamber}.png"
            chart_entry = {
                "order": order,
                "grid_row": (order - 1) // 9,
                "grid_col": (order - 1) % 9,
                "label": f"CH{row.flow_no}-CB{row.chamber_no}",
                "source_file": row.source_file,
                "worksheet": row.chamber,
                "image_path": str(output_path.resolve()),
            }
            group_manifest["charts"].append(chart_entry)
            tasks[row.source_file].append((row, output_path, chart_entry))
        manifest["groups"].append(group_manifest)

    for source_file, source_tasks in tasks.items():
        workbook = load_workbook(base_dir / source_file, read_only=False, data_only=False)
        for row, output_path, chart_entry in source_tasks:
            image_bytes = _target_image(workbook[row.chamber])._data()
            width, height, dpi_x, dpi_y = _quality(image_bytes)
            if width < 2400 or height < 1600 or min(dpi_x, dpi_y) < 299:
                raise RuntimeError(
                    f"Chart quality below 300 dpi target: {source_file}/{row.chamber}: "
                    f"{width}x{height}, dpi=({dpi_x}, {dpi_y})"
                )
            output_path.write_bytes(image_bytes)
            chart_entry.update(
                {
                    "width_px": width,
                    "height_px": height,
                    "dpi_x": dpi_x,
                    "dpi_y": dpi_y,
                }
            )
        workbook.close()

    manifest_path = work_dir / "chart_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest_path, manifest


def _set_text_style(text_frame, font_size: float, bold: bool, color: RGBColor, font_name: str) -> None:
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color


def _add_contained_picture(slide, image_path: Path, left, top, frame_width, frame_height) -> None:
    with Image.open(image_path) as image:
        aspect = image.width / image.height
    frame_aspect = frame_width / frame_height
    if aspect >= frame_aspect:
        width = frame_width
        height = int(frame_width / aspect)
        picture_left = left
        picture_top = top + int((frame_height - height) / 2)
    else:
        height = frame_height
        width = int(frame_height * aspect)
        picture_left = left + int((frame_width - width) / 2)
        picture_top = top
    slide.shapes.add_picture(str(image_path), picture_left, picture_top, width=width, height=height)


def build_growth_curve_ppt(output_path: Path, manifest: dict) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(40)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    margin_x = Inches(0.3125)
    gap_x = Inches(0.125)
    col_width = int((presentation.slide_width - margin_x * 2 - gap_x * 8) / 9)
    image_height = int(col_width / 1.480275)
    title_top = Inches(0.0833)
    title_height = Inches(0.5)
    label_height = Inches(0.2604)
    row_tops = [Inches(0.6667), Inches(3.9792)]

    for group in manifest["groups"]:
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        title_parts = [
            (f"数据组 {group['index']}（n = {group['size']}）", Inches(8.5), "Microsoft YaHei"),
            (f"CV = {group['umax_cv_percent']:.2f}%", Inches(6.5), "Arial"),
            (f"Mean μmax = {group['umax_mean_h-1']:.4f} h⁻¹", Inches(12), "Arial"),
        ]
        title_total_width = sum(part[1] for part in title_parts)
        title_left = int((presentation.slide_width - title_total_width) / 2)
        for title_text, part_width, font_name in title_parts:
            title_box = slide.shapes.add_textbox(
                title_left,
                title_top,
                part_width,
                title_height,
            )
            title_box.text_frame.text = title_text
            _set_text_style(title_box.text_frame, 26, True, RGBColor(17, 24, 39), font_name)
            title_left += part_width

        source_lines = [
            "[Sources]",
            f"- {group['source_workbook']} / {group['selected_sheet']} / μmax mean and sample CV",
        ]
        for chart in group["charts"]:
            left = margin_x + chart["grid_col"] * (col_width + gap_x)
            top = row_tops[chart["grid_row"]]
            label_box = slide.shapes.add_textbox(left, top, col_width, label_height)
            label_box.text_frame.text = chart["label"]
            _set_text_style(label_box.text_frame, 13, True, RGBColor(31, 41, 55), "Arial")
            _add_contained_picture(
                slide,
                Path(chart["image_path"]),
                left,
                top + label_height,
                col_width,
                image_height,
            )
            source_lines.append(
                f"- {chart['source_file']} / {chart['worksheet']} / embedded 300 dpi growth-fitting PNG"
            )

        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is not None:
            notes_frame.text = "\n".join(source_lines)

    presentation.save(output_path)
    return output_path


def compose_group_growth_chart(
    charts: Sequence[dict],
    output_path: Path,
    grid_rows: int = 2,
    grid_cols: int | None = None,
    image_target_width: int = 720,
    image_target_height: int = 480,
    label_height_px: int = 32,
    label_font_pt: int = 16,
    background_rgb: tuple[int, int, int] = (255, 255, 255),
    margin_px: int = 16,
    gap_px: int = 10,
    title: str | None = None,
    title_height_px: int = 36,
    title_font_pt: int = 18,
    dpi: int = 300,
) -> Path:
    """Compose every chamber growth-curve PNG in a group into ONE composite PNG.

    The composite mirrors the PPT slide layout: a centered, multi-column
    ``grid_rows`` × ``grid_cols`` grid where each cell carries one chamber's
    curve with its ``CHx-CBy`` label above it. ``grid_cols`` defaults to
    ``ceil(n / grid_rows)`` so the same 2-row layout used on the slide also
    works for any trailing partial group.

    The output PNG is high-DPI (``dpi`` × ``dpi``) and is meant to be embedded
    once below the data block of the group's ``selected_`` sheet — exactly one
    image per group, regardless of chamber count. Chamber PNGs that are
    missing on disk or fail to decode are skipped silently; an empty input
    raises :class:`RuntimeError` so the caller can decide whether to fail the
    whole workbook write.
    """
    loaded: list[tuple[str, Image.Image]] = []
    for chart in charts:
        image_path = Path(chart["image_path"])
        if not image_path.is_file():
            continue
        try:
            with Image.open(image_path) as source:
                loaded.append((chart.get("label", image_path.stem), source.copy()))
        except (OSError, Image.DecompressionBombError, ValueError):
            continue
    if not loaded:
        raise RuntimeError(f"No usable chamber PNGs found to compose into {output_path}")

    n = len(loaded)
    cols = grid_cols if grid_cols is not None else math.ceil(n / grid_rows)
    rows = grid_rows

    cell_width = image_target_width + 2 * margin_px
    cell_height = (
        label_height_px + image_target_height + 2 * margin_px
    )

    canvas_width = cols * cell_width + max(cols - 1, 0) * gap_px
    canvas_height = rows * cell_height + max(rows - 1, 0) * gap_px
    if title:
        canvas_height += title_height_px + gap_px

    canvas = Image.new("RGB", (canvas_width, canvas_height), background_rgb)
    draw = ImageDraw.Draw(canvas)

    try:
        label_font = ImageFont.truetype("arialbd.ttf", label_font_pt)
    except OSError:
        try:
            label_font = ImageFont.truetype("arial.ttf", label_font_pt)
        except OSError:
            label_font = ImageFont.load_default()
    try:
        title_font = ImageFont.truetype("arialbd.ttf", title_font_pt)
    except OSError:
        title_font = label_font

    if title:
        title_bbox = draw.textbbox((0, 0), title, font=title_font)
        title_w = title_bbox[2] - title_bbox[0]
        title_x = (canvas_width - title_w) // 2
        title_y = (title_height_px - (title_bbox[3] - title_bbox[1])) // 2
        draw.text((title_x, title_y), title, fill=(17, 24, 39), font=title_font)
        grid_origin_y = title_height_px + gap_px
    else:
        grid_origin_y = 0

    for index, (label, image) in enumerate(loaded):
        if index >= cols * rows:
            break
        row, col = divmod(index, cols)
        cell_left = col * (cell_width + gap_px)
        cell_top = grid_origin_y + row * (cell_height + gap_px)

        label_bbox = draw.textbbox((0, 0), label, font=label_font)
        label_w = label_bbox[2] - label_bbox[0]
        label_h = label_bbox[3] - label_bbox[1]
        label_x = cell_left + (cell_width - label_w) // 2
        label_y = cell_top + margin_px
        draw.text((label_x, label_y), label, fill=(17, 24, 39), font=label_font)

        image_area_top = cell_top + margin_px + max(label_h, label_height_px)
        image_box_left = cell_left + margin_px
        image_box_top = image_area_top
        image_box_right = image_box_left + image_target_width
        image_box_bottom = image_box_top + image_target_height

        aspect = image.width / image.height
        target_aspect = image_target_width / image_target_height
        if aspect >= target_aspect:
            new_width = image_target_width
            new_height = max(1, int(image_target_width / aspect))
        else:
            new_height = image_target_height
            new_width = max(1, int(image_target_height * aspect))
        resized = image.resize((new_width, new_height), Image.LANCZOS)

        paste_left = image_box_left + (image_target_width - new_width) // 2
        paste_top = image_box_top + (image_target_height - new_height) // 2
        canvas.paste(resized, (paste_left, paste_top))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, dpi=(dpi, dpi))
    return output_path

from __future__ import annotations

import json
import re
import shutil
import zipfile
from copy import copy
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl
from openpyxl import load_workbook
from openpyxl.drawing.image import Image

from .grouping import metrics, sort_for_output
from .models import ChamberResult
from .ppt_io import compose_group_growth_chart


SUMMARY_SHEET = "汇总结果"
CHAMBER_PATTERN = re.compile(r"^CH(\d+)_CB(\d+)$")
OUTPUT_HEADERS = 13


def _number(value: object, source: Path, row_number: int, column_number: int) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Non-numeric value in {source.name}, row {row_number}, column {column_number}: {value!r}"
        ) from exc


def discover_source_workbooks(base_dir: Path) -> list[Path]:
    sources: list[Path] = []
    for path in sorted(base_dir.glob("*.xlsx"), key=lambda item: item.name):
        name = path.name
        if name.endswith("_marked.xlsx") or name.startswith("selected__groups_"):
            continue
        if name == "参考输出.xlsx":
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            is_source = SUMMARY_SHEET in workbook.sheetnames
            workbook.close()
        except (OSError, KeyError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException):
            is_source = False
        if is_source:
            sources.append(path)
    if not sources:
        raise RuntimeError(f"No source workbooks containing '{SUMMARY_SHEET}' were found in {base_dir}")
    return sources


def collect_candidates(source_paths: Sequence[Path]) -> list[ChamberResult]:
    rows: list[ChamberResult] = []
    for source in source_paths:
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook[SUMMARY_SHEET]
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            chamber = values[0] if values else None
            match = CHAMBER_PATTERN.fullmatch(chamber) if isinstance(chamber, str) else None
            if not match:
                continue
            numeric = [
                _number(values[index] if index < len(values) else None, source, row_number, index + 1)
                for index in range(1, 11)
            ]
            n0, r2 = numeric[2], numeric[3]
            if n0 <= 0 or r2 <= 0.9:
                continue
            flow_no, chamber_no = int(match.group(1)), int(match.group(2))
            rows.append(
                ChamberResult(
                    source_file=source.name,
                    summary_row=row_number,
                    chamber=chamber,
                    flow=f"CH{flow_no}",
                    flow_no=flow_no,
                    chamber_no=chamber_no,
                    umax=numeric[0],
                    k=numeric[1],
                    n0=n0,
                    r2=r2,
                    td=numeric[4],
                    f=numeric[5],
                    eta=numeric[6],
                    lag=numeric[7],
                    log=numeric[8],
                    stable=numeric[9],
                )
            )
        workbook.close()
    return rows


def _copy_row_style(sheet, source_row: int, target_row: int, columns: int) -> None:
    if source_row == target_row:
        return
    for column in range(1, columns + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        if source.has_style:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _clear_values(sheet, start_row: int = 2, end_row: int = 200, columns: int = OUTPUT_HEADERS) -> None:
    for row in sheet.iter_rows(min_row=start_row, max_row=max(end_row, sheet.max_row), min_col=1, max_col=columns):
        for cell in row:
            cell.value = None


def _format_output_sheet(sheet, end_row: int) -> None:
    widths = {
        "A": 30,
        "B": 16,
        "C": 16,
        "D": 23,
        "E": 22,
        "F": 22,
        "G": 16,
        "H": 20,
        "I": 16,
        "J": 22,
        "K": 16,
        "L": 16,
        "M": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 42
    for cell in sheet[1]:
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment
    for row_number in range(2, end_row + 1):
        sheet.row_dimensions[row_number].height = 21
        sheet.cell(row_number, 4).number_format = "0.0000"
        sheet.cell(row_number, 5).number_format = "0.00"
        sheet.cell(row_number, 6).number_format = "0.00"
        sheet.cell(row_number, 7).number_format = "0.0000"
        sheet.cell(row_number, 8).number_format = "0.00"
        sheet.cell(row_number, 9).number_format = "0.000"
        sheet.cell(row_number, 10).number_format = "0.000"
        for column in range(11, 14):
            sheet.cell(row_number, column).number_format = "0.00"
    sheet.freeze_panes = "A2"


def _write_table(sheet, rows: Iterable[ChamberResult]) -> int:
    ordered = sort_for_output(rows)
    for row_number, result in enumerate(ordered, start=2):
        _copy_row_style(sheet, 2, row_number, OUTPUT_HEADERS)
        for column, value in enumerate(result.output_row(), start=1):
            sheet.cell(row_number, column).value = value
    end_row = len(ordered) + 1
    _format_output_sheet(sheet, end_row)
    return end_row


def build_group_workbooks(
    reference_path: Path,
    output_dir: Path,
    groups: Sequence[Sequence[ChamberResult]],
    strict_rows: Sequence[ChamberResult],
    source_paths: Sequence[Path],
    min_group_size: int,
    max_group_size: int,
) -> list[Path]:
    outputs: list[Path] = []
    full_group_count = sum(len(group) == max_group_size for group in groups)
    sub_group_sizes = (
        ", ".join(str(len(group)) for group in groups if len(group) < max_group_size) or "None"
    )

    for index, group in enumerate(groups, start=1):
        workbook = load_workbook(reference_path, data_only=False)
        selected_sheet = workbook["selected_18"]
        strict_sheet = workbook["strict_19"]
        summary_sheet = workbook["summary"]
        selected_sheet.title = f"selected_{len(group)}"
        strict_sheet.title = f"strict_{len(strict_rows)}"

        selected_style = [copy(selected_sheet.cell(2, col)._style) for col in range(1, OUTPUT_HEADERS + 1)]
        strict_style = [copy(strict_sheet.cell(2, col)._style) for col in range(1, OUTPUT_HEADERS + 1)]
        summary_style = [copy(summary_sheet.cell(2, col)._style) for col in range(1, 3)]
        _clear_values(selected_sheet)
        _clear_values(strict_sheet)
        _clear_values(summary_sheet, columns=2)
        for col, style in enumerate(selected_style, start=1):
            selected_sheet.cell(2, col)._style = copy(style)
        for col, style in enumerate(strict_style, start=1):
            strict_sheet.cell(2, col)._style = copy(style)
        for col, style in enumerate(summary_style, start=1):
            summary_sheet.cell(2, col)._style = copy(style)

        selected_end = _write_table(selected_sheet, group)
        strict_end = _write_table(strict_sheet, strict_rows)
        if selected_sheet.max_row > selected_end:
            selected_sheet.delete_rows(selected_end + 1, selected_sheet.max_row - selected_end)
        if strict_sheet.max_row > strict_end:
            strict_sheet.delete_rows(strict_end + 1, strict_sheet.max_row - strict_end)

        result = metrics(group)
        summary_rows = [
            ("source_files", "; ".join(path.name for path in source_paths)),
            ("filter_rule", "N0 > 0, R2 > 0.9, CV < 15%, JB p > 0.05"),
            ("strict_candidate_total", len(strict_rows)),
            ("group_index", index),
            ("group_size", len(group)),
            ("target_full_group_size", max_group_size),
            ("minimum_group_size", min_group_size),
            ("full_target_group_count", full_group_count),
            ("sub_target_group_sizes", sub_group_sizes),
            ("mu_max_mean_h^-1", result.mean),
            ("mu_max_sample_sd", result.sample_sd),
            ("mu_max_cv_percent", result.cv * 100),
            ("jarque_bera_statistic", result.jb),
            ("jarque_bera_p_value", result.p_value),
            ("ordering", "Flow ascending, then chamber number ascending"),
        ]
        for row_number, values in enumerate(summary_rows, start=2):
            _copy_row_style(summary_sheet, 2, row_number, 2)
            summary_sheet.cell(row_number, 1).value = values[0]
            summary_sheet.cell(row_number, 2).value = values[1]
            if values[0] in {
                "mu_max_mean_h^-1",
                "mu_max_sample_sd",
                "mu_max_cv_percent",
                "jarque_bera_statistic",
                "jarque_bera_p_value",
            }:
                summary_sheet.cell(row_number, 2).number_format = "0.000000"
        summary_sheet.freeze_panes = "A2"
        if summary_sheet.max_row > len(summary_rows) + 1:
            summary_sheet.delete_rows(len(summary_rows) + 2, summary_sheet.max_row - len(summary_rows) - 1)

        output_path = output_dir / f"selected__groups_{index}_{len(group)}.xlsx"
        workbook.save(output_path)
        workbook.close()
        with zipfile.ZipFile(output_path) as archive:
            if archive.testzip() is not None:
                raise RuntimeError(f"Corrupt workbook written: {output_path}")
        outputs.append(output_path)
    return outputs


def _set_font_red(cell) -> None:
    font = copy(cell.font)
    font.color = "FFFF0000"
    cell.font = font


def mark_source_workbooks(
    base_dir: Path,
    output_dir: Path,
    source_paths: Sequence[Path],
    groups: Sequence[Sequence[ChamberResult]],
) -> list[Path]:
    selected_by_file: dict[str, list[ChamberResult]] = {path.name: [] for path in source_paths}
    for group in groups:
        for row in group:
            selected_by_file[row.source_file].append(row)

    outputs: list[Path] = []
    for source in source_paths:
        workbook = load_workbook(base_dir / source.name, read_only=False, data_only=False)
        summary = workbook[SUMMARY_SHEET]
        for result in selected_by_file[source.name]:
            for column in range(1, summary.max_column + 1):
                _set_font_red(summary.cell(result.summary_row, column))
            chamber_sheet = workbook[result.chamber]
            for row in chamber_sheet.iter_rows(
                min_row=1,
                max_row=chamber_sheet.max_row,
                min_col=1,
                max_col=chamber_sheet.max_column,
            ):
                for cell in row:
                    _set_font_red(cell)

        output_path = output_dir / f"{source.stem}_marked.xlsx"
        workbook.save(output_path)
        workbook.close()
        with zipfile.ZipFile(output_path) as archive:
            if archive.testzip() is not None:
                raise RuntimeError(f"Corrupt marked workbook written: {output_path}")
        outputs.append(output_path)
    return outputs


def write_partition_results(
    path: Path,
    source_paths: Sequence[Path],
    strict_rows: Sequence[ChamberResult],
    groups: Sequence[Sequence[ChamberResult]],
    remaining: Sequence[ChamberResult],
    window_count: int,
    min_group_size: int,
    max_group_size: int,
    cv_limit: float,
    normality_p_limit: float,
    mean_spread_ratio_limit: float | None = None,
) -> None:
    payload = {
        "sourceFiles": [source.name for source in source_paths],
        "strictCount": len(strict_rows),
        "windowCount": window_count,
        "groupRules": {
            "minimumGroupSize": min_group_size,
            "preferredFullGroupSize": max_group_size,
            "cvLimit": cv_limit,
            "normalityPValueLimit": normality_p_limit,
            "maxMeanSpreadRatio": mean_spread_ratio_limit,
        },
        "groups": [
            {
                "index": index,
                "size": len(group),
                "stats": metrics(group).to_dict(),
                "members": [
                    {"file": row.source_file, "chamber": row.chamber, "umax": row.umax}
                    for row in sort_for_output(group)
                ],
            }
            for index, group in enumerate(groups, start=1)
        ],
        "remainingStrictCandidates": [
            {"file": row.source_file, "chamber": row.chamber, "umax": row.umax}
            for row in sort_for_output(remaining)
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_umax_summary(
    path: Path,
    groups: Sequence[Sequence[ChamberResult]],
) -> Path:
    """Write a per-group μmax companion workbook: one row per group.

    Each row holds the group's *mean* μmax (the same headline value shown on
    the generated PPT slide for that group), so users can quickly look up
    the headline growth-rate metric for every selected batch without paging
    through the per-group workbooks. This summary is a flat companion and
    does not change any of the existing per-group workbooks.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "group_mean_umax"
    sheet.cell(1, 1).value = "Group"
    sheet.cell(1, 2).value = "Mean μmax (h⁻¹)"
    for index, group in enumerate(groups, start=1):
        sheet.cell(index + 1, 1).value = index
        sheet.cell(index + 1, 2).value = metrics(group).mean
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 20
    sheet.freeze_panes = "A2"
    workbook.save(path)
    workbook.close()
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError(f"Corrupt umax-summary workbook written: {path}")
    return path


def embed_chamber_images(
    group_workbook_path: Path,
    charts: Sequence[dict],
    embedded_width_px: int = 1600,
) -> None:
    """Embed ONE composite growth-curve PNG covering the entire group.

    All chambers in the group are composited by
    :func:`ppt_io.compose_group_growth_chart` into a single labeled
    multi-cell image mirroring the PPT slide layout (2-row grid). The
    composite is saved next to the per-chamber PNGs already extracted for
    the PPT slide (``work_dir/ppt_charts/group_i/composite.png``) and then
    anchored in the group's ``selected_`` sheet directly below the data
    block — exactly one image per group, regardless of chamber count.

    ``embedded_width_px`` controls the displayed width in the workbook;
    the height follows the composite's intrinsic aspect ratio so all
    chamber growth curves remain legible. Row height is scaled to keep the
    composite fully visible without truncating the chamber data above it.
    """
    workbook = load_workbook(group_workbook_path)
    sheet_name = next(
        name for name in workbook.sheetnames if name.startswith("selected_")
    )
    sheet = workbook[sheet_name]
    n_data_rows = len(charts)
    first_image_row = 2 + n_data_rows  # immediately below the last data row

    if not charts:
        workbook.save(group_workbook_path)
        workbook.close()
        return

    composite_path = (
        Path(charts[0]["image_path"]).parent / "composite.png"
    )
    group_index = composite_path.parent.name  # e.g. "group_1"
    group_headline = (
        f"Data Group {group_index.replace('group_', '')} (n = {n_data_rows})"
    )

    compose_group_growth_chart(
        charts,
        composite_path,
        title=group_headline,
        grid_rows=2,
    )

    if not composite_path.is_file():
        workbook.save(group_workbook_path)
        workbook.close()
        return

    from PIL import Image as PILImage
    with PILImage.open(composite_path) as probe:
        canvas_w, canvas_h = probe.size

    display_width = embedded_width_px
    display_height = max(1, int(display_width * canvas_h / canvas_w))
    row_height_pt = max(180, int(display_height / 96 * 72) + 12)

    img = Image(str(composite_path))
    img.width = display_width
    img.height = display_height
    img.anchor = f"A{first_image_row}"
    sheet.add_image(img)
    sheet.row_dimensions[first_image_row].height = row_height_pt

    workbook.save(group_workbook_path)
    workbook.close()


def clean_generated_outputs(output_dir: Path) -> None:
    for path in output_dir.glob("selected__groups_*_*.xlsx"):
        path.unlink()
    for path in output_dir.glob("*_marked.xlsx"):
        path.unlink()
    pptx = output_dir / "selected_groups_growth_curves_2x9.pptx"
    if pptx.exists():
        pptx.unlink()
    summary = output_dir / "selected_groups_umax_summary.xlsx"
    if summary.exists():
        summary.unlink()


def handle_root_ndjson(output_dir: Path, logs_dir: Path, policy: str) -> None:
    root_logs = list(output_dir.glob("*.ndjson"))
    if policy == "keep":
        return
    if policy == "delete":
        for path in root_logs:
            path.unlink()
        return
    logs_dir.mkdir(parents=True, exist_ok=True)
    for path in root_logs:
        destination = logs_dir / path.name
        if destination.exists():
            destination.unlink()
        shutil.move(str(path), str(destination))

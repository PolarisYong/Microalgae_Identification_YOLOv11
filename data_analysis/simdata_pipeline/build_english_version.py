# -*- coding: utf-8 -*-
"""
入口⑥：英文版 21 组数据（Group21_English_Version）
==================================================
把「20组模拟数据+1组最优模拟」里的每组三个文件（_可视化结果.xlsx /
.pptx / selected__groups_*.xlsx）整体转为英文版，输出到
``模拟数据处理_part1/Group21_English_Version/<同名子文件夹>/``。

转换口径（均已与用户确认）：
- 模板：可视化结果参考模板.xlsx —— 「汇总结果」页（8 列腔室表 + 培养条件块 +
  Pooled fitting results + 3 张合并图）、CB-1..CB-n 腔室页（A:F + 10 参数 + 4 图）
- 图表文字：按「图表文字内容修改意见.pptx」——全部生长曲线改 Y=ln(Nrel) 尺度
  （散点 ln(实测)、曲线按 Y(t) 公式用已存参数解析生成，不重拟合），
  参数框 λ 带单位 h、R² 写作 R²_Y，标题/坐标轴用 DIC 条件标签
- Tlag = λ（可为负）、Tprolif = t90 − λ（与修改意见严格一致）
- 页签 CB-1..CB-n；A 列图片名 CB{k}_H{i}.tif；状态 成功→Success
- PPT 页眉「数据组 i（n = N）→ Group i (n = N)」，标签 → CB-<k>，
  配图按原像素尺寸重绘替换
- selected__groups：selected_/strict_ 页表头中→英，数据原样

文件夹内 xlsx↔selected↔pptx 的对应关系按**原始腔室集合**匹配（数据来源页 /
selected 页腔室名称列 / PPT 标签三方取交集相等的）；数量不一致时在日志中
明确列出，缺项跳过、多余项不处理。2208 的 数据溯源.xlsx 原样复制。

用法：
    python build_english_version.py            # 全量（已存在的输出跳过）
    python build_english_version.py --force    # 覆盖重算
    python build_english_version.py --dry-run  # 只打印计划，不写文件
    python build_english_version.py --only "L30‑N20‑IC0.5%"  # 只转一个文件夹
"""
import argparse
import re
import shutil
import sys
from pathlib import Path

import config
from core import en_combo, en_labels, en_pptx, en_selected, en_writer
from core.logging_setup import setup_logger

SRC_ROOT = config.PART1_ROOT / "20组模拟数据+1组最优模拟"
DST_ROOT = config.PART1_ROOT / "Group21_English_Version"

SUMMARY_SHEET = "汇总"
SOURCE_SHEET = "数据来源"
MERGED_TITLE = "合并拟合结果"

XLSX_RE = re.compile(r"_可视化结果(_\d+)?\.xlsx$")
SELECTED_RE = re.compile(r"^selected__groups_\d+_\d+\.xlsx$")
PPT_LABEL_SRC_RE = re.compile(r"^CH(\d+)-CB(\d+)$")
SOURCE_SHEET_RE = re.compile(r"页签：\s*(CH\d+_CB\d+)\s*$")
CHAMBER_NAME_RE = re.compile(r"^(?:CH\d+_)?CB(\d+)$")


# --------------------------------------------------------------------- 读源工作簿
def read_group_xlsx(path, logger, tag):
    """
    读一个 _可视化结果.xlsx → dict(chambers, merged, src_set, src_to_index, suffix)

    chambers: [{index, cb_id, data_rows, counts, areas, params, t_end}]
    merged  : 合并拟合参数（规范键）或 None
    src_set : 原始腔室名集合（无数据来源页签时为 None）
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        # 腔室顺序：汇总页 A 列（CH1_CBn 或 CBn）
        order = []
        if SUMMARY_SHEET in wb.sheetnames:
            for row in wb[SUMMARY_SHEET].iter_rows(min_row=2, values_only=True):
                value = row[0] if row else None
                if isinstance(value, str) and CHAMBER_NAME_RE.fullmatch(value.strip()):
                    order.append(value.strip())
                elif order:
                    break
        if not order:
            order = [s for s in wb.sheetnames
                     if s not in (SUMMARY_SHEET, SOURCE_SHEET)
                     and CHAMBER_NAME_RE.fullmatch(s)]
        if not order:
            raise ValueError(f"{path.name} 里找不到腔室页签")

        chambers = []
        for index, sheet in enumerate(order, 1):
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            body = rows[1:]
            data_rows, counts, areas, params_cn = [], [], [], {}
            for r in body:
                status = r[1] if len(r) > 1 else None
                count = r[config.COUNTS_COLUMN_INDEX] \
                    if len(r) > config.COUNTS_COLUMN_INDEX else None
                area = r[config.AREA_COLUMN_INDEX] \
                    if len(r) > config.AREA_COLUMN_INDEX else None
                data_rows.append((status, count, area))
                if count is not None:
                    counts.append(count)
                if area is not None:
                    areas.append(area)
                if len(r) > 5 and r[4] is not None:
                    key = str(r[4]).strip()
                    if key and key not in ("参数名", "Parameter"):
                        params_cn[key] = r[5]
            chambers.append({
                "index": index,
                "cb_id": f"CB-{index}",
                "src_sheet": sheet,
                "data_rows": data_rows,
                "counts": counts,
                "areas": areas,
                "params": en_labels.map_params(params_cn),
                "t_end": len(counts) - 1,
            })

        # 合并拟合结果块
        merged = None
        if SUMMARY_SHEET in wb.sheetnames:
            block, found = {}, False
            for row in wb[SUMMARY_SHEET].iter_rows(values_only=True):
                label = row[0] if row else None
                if label == MERGED_TITLE:
                    found = True
                    continue
                if found and isinstance(label, str) and label.strip():
                    block[label.strip()] = row[1] if len(row) > 1 else None
            if block:
                merged = en_labels.map_params(block, en_labels.CN_MERGED_KEYS)
                merged["n_worksheets"] = block.get("参与拟合的页签数量")
                merged["n_points"] = block.get("总数据点数量")

        # 数据来源：腔室序号 ↔ 原始腔室名
        # 注意：同一原始腔室可能被重复纳入一组（如 CH1_CB23 出现两次），
        # src_to_index 只保留最后出现（旧行为），src_to_indices 保留全部
        # （按出现顺序），PPT 标签 / 组合图按出现顺序消费。
        src_set, src_to_index, src_to_indices = None, {}, {}
        if SOURCE_SHEET in wb.sheetnames:
            src_set = set()
            for row in wb[SOURCE_SHEET].iter_rows(min_row=2, values_only=True):
                name, source = (row + (None, None))[:2] if row else (None, None)
                m = CHAMBER_NAME_RE.fullmatch(str(name).strip()) if name else None
                s = SOURCE_SHEET_RE.search(str(source)) if source else None
                if m and s:
                    src_set.add(s.group(1))
                    src_to_index[s.group(1)] = int(m.group(1))
                    src_to_indices.setdefault(s.group(1), []).append(int(m.group(1)))

        suffix = XLSX_RE.search(path.name)
        return {
            "chambers": chambers,
            "merged": merged,
            "src_set": src_set,
            "src_to_index": src_to_index,
            "src_to_indices": src_to_indices,
            "suffix": suffix.group(1) or "" if suffix else "",
            "path": path,
        }
    finally:
        wb.close()


def read_selected_src_set(path):
    """selected__groups 文件 → selected_ 页腔室名称集合（原始腔室名）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        name = next((n for n in wb.sheetnames if n.startswith("selected_")), None)
        if name is None:
            return None
        ws = wb[name]
        rows = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        try:
            col = next(i for i, v in enumerate(rows)
                       if isinstance(v, str) and "腔室名称" in v)
        except StopIteration:
            return None
        out = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) > col and isinstance(row[col], str):
                out.add(row[col].strip())
        return out or None
    finally:
        wb.close()


def read_pptx_src_set(path):
    """PPT 标签 → 原始腔室名集合；标签为 CB 序号时返回空集合。"""
    from pptx import Presentation

    prs = Presentation(str(path))
    out = set()
    for sh in prs.slides[0].shapes:
        if not sh.has_text_frame:
            continue
        m = PPT_LABEL_SRC_RE.fullmatch(sh.text_frame.text.strip())
        if m:
            out.add(f"CH{m.group(1)}_CB{m.group(2)}")
    return out


# --------------------------------------------------------------------- 匹配
def match_assets(groups, selected_files, pptx_files, logger, tag):
    """
    按原始腔室集合把 selected / pptx 分配给各组。
    返回 {group_index: {"selected": path|None, "pptx": path|None}}。
    """
    plan = {i: {"selected": None, "pptx": None} for i in range(len(groups))}
    used_s, used_p = set(), set()

    for gi, g in enumerate(groups):
        src_set = g["src_set"]
        if src_set is None:                         # 无数据来源：仅在单组文件夹可配
            if len(groups) == 1:
                if len(selected_files) == 1:
                    plan[gi]["selected"] = selected_files[0]
                    used_s.add(selected_files[0])
                if len(pptx_files) == 1:
                    plan[gi]["pptx"] = pptx_files[0]
                    used_p.add(pptx_files[0])
            continue
        for p in selected_files:
            if p in used_s:
                continue
            if read_selected_src_set(p) == src_set:
                plan[gi]["selected"] = p
                used_s.add(p)
                break
        for p in pptx_files:
            if p in used_p:
                continue
            ppt_set = read_pptx_src_set(p)
            if ppt_set and ppt_set == src_set:
                plan[gi]["pptx"] = p
                used_p.add(p)
                break
        if plan[gi]["selected"] is None:
            logger.warning("%s %s 未匹配到 selected__groups 文件", tag, g["path"].name)
        if plan[gi]["pptx"] is None:
            logger.warning("%s %s 未匹配到 pptx 文件", tag, g["path"].name)

    for p in set(selected_files) - used_s:
        logger.warning("%s selected 文件 %s 不属于本文件夹任何组，跳过", tag, p.name)
    for p in set(pptx_files) - used_p:
        logger.warning("%s pptx 文件 %s 不属于本文件夹任何组（可能误放），跳过",
                       tag, p.name)
    return plan


# --------------------------------------------------------------------- 单组转换
def convert_group(group, assets, cond, out_dir, logger, tag, force=False,
                  dry_run=False):
    """转换一组的三个文件，返回产物路径列表。"""
    suffix = group["suffix"]
    xlsx_out = out_dir / f"{cond['label']}_visualization_results{suffix}.xlsx"
    pptx_out = out_dir / f"{cond['label']}{suffix}.pptx"
    outputs = []

    todo = []
    if xlsx_out.exists() and not force:
        logger.info("  %s 已存在，跳过：%s", tag, xlsx_out.name)
    else:
        todo.append("xlsx")
    if assets["pptx"] is not None and (not pptx_out.exists() or force):
        todo.append("pptx")
    sel_out = None
    if assets["selected"] is not None:
        sel_out = out_dir / assets["selected"].name
        if not sel_out.exists() or force:
            todo.append("selected")

    if dry_run:
        logger.info("  %s [dry-run] 将生成：%s%s%s", tag,
                    xlsx_out.name,
                    f"、{pptx_out.name}" if assets["pptx"] else "（无 pptx）",
                    f"、{assets['selected'].name}" if assets["selected"] else "")
        return [xlsx_out, pptx_out, sel_out]

    out_dir.mkdir(parents=True, exist_ok=True)
    if "xlsx" in todo:
        en_writer.build_english_workbook(
            xlsx_out, group["chambers"], group["merged"], cond,
            logger=logger, tag=tag)
    outputs.append(xlsx_out)
    if "pptx" in todo:
        en_pptx.rebuild_pptx_en(
            assets["pptx"], pptx_out, group["chambers"], cond,
            group.get("src_to_indices") or group["src_to_index"],
            logger=logger, tag=tag)
        outputs.append(pptx_out)
    if "selected" in todo:
        en_selected.translate_selected(assets["selected"], sel_out,
                                       logger=logger, tag=tag)
        # 组合生长曲线图重绘（英文标题/Y 尺度/Times 字体，原位替换）
        try:
            en_combo.rebuild_combo(sel_out, group["chambers"], cond,
                                   group.get("src_to_indices")
                                   or group["src_to_index"],
                                   logger=logger, tag=tag)
        except Exception as exc:                            # noqa: BLE001
            logger.error("%s 组合图重绘失败：%s（selected 保留原图）", tag, exc)
        outputs.append(sel_out)
    return outputs


# --------------------------------------------------------------------- 主流程
def main(argv=None):
    parser = argparse.ArgumentParser(description="英文版 21 组数据批处理")
    parser.add_argument("--dry-run", action="store_true", help="只打印计划")
    parser.add_argument("--force", action="store_true", help="覆盖已存在的输出")
    parser.add_argument("--only", default=None, help="只处理指定文件夹名")
    args = parser.parse_args(argv)

    logger = setup_logger(config.PROJECT_ROOT / "simdata_pipeline" / "_logs",
                          name="build_english")
    logger.info("=" * 70)
    logger.info("英文版转换：%s → %s", SRC_ROOT, DST_ROOT)
    logger.info("模式：%s%s", "dry-run " if args.dry_run else "",
                "force" if args.force else "跳过已存在")

    if not SRC_ROOT.is_dir():
        logger.error("源目录不存在：%s", SRC_ROOT)
        return 1

    stats = {"folders": 0, "groups": 0, "xlsx": 0, "pptx": 0, "selected": 0,
             "skipped_folders": []}
    for folder in sorted(p for p in SRC_ROOT.iterdir() if p.is_dir()):
        if args.only and folder.name != args.only:
            continue
        tag = f"[{folder.name}]"
        cond = en_labels.condition_of_folder(folder.name)
        if cond is None:
            logger.warning("%s 无法解析培养条件，整个文件夹跳过", tag)
            stats["skipped_folders"].append(folder.name)
            continue
        stats["folders"] += 1
        logger.info("%s 条件 %s coded %s", tag, cond["label"], cond["coded_text"])

        group_files = sorted(p for p in folder.glob("*.xlsx")
                             if XLSX_RE.search(p.name)
                             and not p.name.startswith("~$"))
        selected_files = sorted(p for p in folder.glob("*.xlsx")
                                if SELECTED_RE.match(p.name))
        pptx_files = sorted(p for p in folder.glob("*.pptx")
                            if not p.name.startswith("~$"))
        if not group_files:
            logger.warning("%s 没有 _可视化结果 xlsx，跳过", tag)
            continue

        groups = []
        for p in group_files:
            try:
                groups.append(read_group_xlsx(p, logger, tag))
            except Exception as exc:                            # noqa: BLE001
                logger.error("%s 读取 %s 失败：%s", tag, p.name, exc)
        if not groups:
            continue
        stats["groups"] += len(groups)

        plan = match_assets(groups, selected_files, pptx_files, logger, tag)
        out_dir = DST_ROOT / folder.name
        for gi, g in enumerate(groups):
            gtag = f"{tag}{g['path'].name}"
            outs = convert_group(g, plan[gi], cond, out_dir, logger, gtag,
                                 force=args.force, dry_run=args.dry_run)
            for o in outs:
                if o is None:
                    continue
                key = ("xlsx" if o.name.endswith("_可视化结果.xlsx")
                       or "visualization_results" in o.name
                       else "pptx" if o.suffix == ".pptx" else "selected")
                stats[key] += 1

        # 其余文件（如 2208 的 数据溯源.xlsx）原样复制
        handled = set(group_files) | set(selected_files)
        for p in folder.glob("*.xlsx"):
            if p in handled or p.name.startswith("~$"):
                continue
            dst = out_dir / p.name
            if args.dry_run:
                logger.info("  %s [dry-run] 原样复制：%s", tag, p.name)
            elif not dst.exists() or args.force:
                out_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(p, dst)
                logger.info("  %s 原样复制：%s", tag, p.name)

    logger.info("=" * 70)
    logger.info("完成：文件夹 %d 个，数据组 %d 个；生成 工作簿 %d / PPT %d / "
                "selected %d", stats["folders"], stats["groups"],
                stats["xlsx"], stats["pptx"], stats["selected"])
    if stats["skipped_folders"]:
        logger.warning("跳过的文件夹：%s", ", ".join(stats["skipped_folders"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())

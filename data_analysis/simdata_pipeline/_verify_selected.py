# -*- coding: utf-8 -*-
"""批量校验所有 selected__groups 产物。"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import openpyxl

DST = Path(r"E:/pythonProject/Microalgae_Identification_YOLOv11/data_analysis/模拟数据处理_part1/Group21_English_Version")
MAP = Path(r"E:/pythonProject/Microalgae_Identification_YOLOv11/data_analysis/simdata_pipeline/_logs/selected_summary_pairs.json")

mapping = json.loads(MAP.read_text(encoding="utf-8"))
print(f"配对清单 {len(mapping)} 条")

ok = bad = 0
for key in sorted(mapping):
    folder_name, sel_name = key.split("/", 1)
    viz_name = mapping[key]
    sel = DST / folder_name / sel_name
    viz = DST / folder_name / viz_name
    problems = []

    wb = openpyxl.load_workbook(sel, rich_text=True)
    if wb.sheetnames != [sel.stem.split("__groups_")[0] + "_" +
                         sel_name.split("_")[-1].replace(".xlsx", "")]:
        pass
    if len(wb.sheetnames) != 1:
        problems.append(f"页签数={len(wb.sheetnames)} {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]

    src = openpyxl.load_workbook(viz, rich_text=True)["汇总结果"]
    # 找条件块末行
    bottom = None
    for r in range(1, src.max_row + 1):
        v = src.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("Cultivation conditions"):
            bottom = r + 2
            break
    if bottom is None:
        problems.append("源无条件块")
    else:
        diff = 0
        for r in range(1, bottom + 1):
            for c in range(1, 9):
                if src.cell(row=r, column=c).value != ws.cell(row=r, column=c).value:
                    diff += 1
        if diff:
            problems.append(f"值差异 {diff} 格")
        # 图片锚点
        rows = [im.anchor._from.row for im in ws._images]
        if rows != [bottom + 1] * len(rows):
            problems.append(f"图片锚点 {rows} 应为 {bottom + 1}")
        if ws.row_dimensions[bottom + 2].height != 212:
            problems.append(f"图行高 {ws.row_dimensions[bottom + 2].height}")
        # 残留
        left = sum(1 for r in range(1, ws.max_row + 1) for c in range(9, 14)
                   if ws.cell(row=r, column=c).value is not None)
        if left:
            problems.append(f"I-M 残留 {left}")
        if ws["A1"].value != "Valid cultivation unit ID":
            problems.append(f"A1={ws['A1'].value!r}")
        if ws.cell(row=bottom, column=1).value in (None, ""):
            problems.append("条件取值缺失")
    wb.close()

    if problems:
        bad += 1
        print(f"✗ {folder_name}/{sel_name}\n    " + "; ".join(problems))
    else:
        ok += 1
        print(f"✓ {folder_name}/{sel_name}  ←  {viz_name}  (末行 {bottom})")

print(f"\n通过 {ok} / 失败 {bad}")

# -*- coding: utf-8 -*-
"""
入口⑦：把可视化工作簿「汇总结果」页的培养条件区并入 selected__groups 文件
========================================================================
仅处理英文版 ``模拟数据处理_part1/Group21_English_Version/``（中文版不动）。

对每个 ``selected__groups_X_Y.xlsx``：
1. 找到它在同一文件夹里对应的 ``<标签>_visualization_results<suffix>.xlsx``
   （按**原始腔室集合**匹配，与入口⑥同一套逻辑：读源目录
   ``20组模拟数据+1组最优模拟/<同名文件夹>/`` 的「数据来源」页取集合）；
2. 把该可视化工作簿「汇总结果」页 ``A1:H<培养条件取值行>`` 整块复制进
   ``selected_Y`` 页，替换原来的 13 列腔室表（含 Cultivation conditions
   条件块，如截图所示）；
3. 把 ``selected_Y`` 页嵌的组合生长曲线图整体下移到数据块下方（不压数据）；
4. 删除 ``selected_Y`` 以外的页签（strict_M / summary）。

输出：原地覆盖 selected__groups 文件；改前整份备份到
``_backup/selected_cond_<时间戳>/``。

用法：
    python link_selected_summary.py --dry-run     # 只打印计划
    python link_selected_summary.py               # 执行（已处理过的跳过）
    python link_selected_summary.py --force       # 强制重做
    python link_selected_summary.py --only "L120‑N20‑IC5.25%"
"""
import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import config
from core import selected_summary
from core.logging_setup import setup_logger

import build_english_version as bev

SRC_ROOT = config.PART1_ROOT / "20组模拟数据+1组最优模拟"
DST_ROOT = config.PART1_ROOT / "Group21_English_Version"
BACKUP_ROOT = config.PROJECT_ROOT / "_backup"
MAP_FILE = config.PROJECT_ROOT / "simdata_pipeline" / "_logs" / \
    "selected_summary_pairs.json"


# --------------------------------------------------------------------- 配对清单
def load_map():
    try:
        return json.loads(MAP_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_map(mapping):
    MAP_FILE.parent.mkdir(parents=True, exist_ok=True)
    MAP_FILE.write_text(json.dumps(mapping, ensure_ascii=False, indent=2,
                                   sort_keys=True), encoding="utf-8")


# --------------------------------------------------------------------- 读 selected 腔室列
def read_selected_chambers(path):
    """
    英文版 selected_ 页 → (页签名, [原始腔室名, ...])。
    表头兼容英文 'Chamber' 与中文 '腔室名称'；已处理过的页签（表头变成
    'Valid cultivation unit ID'）返回空列表。
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        name = next((n for n in wb.sheetnames if n.startswith("selected_")), None)
        if name is None:
            return None, []
        ws = wb[name]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col = next((i for i, h in enumerate(header)
                    if isinstance(h, str)
                    and (h.strip().lower() == "chamber" or "腔室名称" in h)),
                   None)
        if col is None:
            return name, []
        names = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = row[col] if len(row) > col else None
            if isinstance(v, str) and v.strip():
                names.append(v.strip())
        return name, names
    finally:
        wb.close()


def chamber_set_of(folder, sel_path):
    """
    EN selected 文件的腔室集合；已被处理过（腔室列被覆盖）时回退读源目录
    同名 selected__groups 文件（中文表头，含原始腔室名）。
    """
    _, names = read_selected_chambers(sel_path)
    if names:
        return set(names)
    src = SRC_ROOT / folder.name / sel_path.name
    if src.is_file():
        _, names = read_selected_chambers(src)
        return set(names)
    return set()


def is_done(path):
    """已经只剩一个页签 = 处理过（幂等跳过）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return len(wb.sheetnames) <= 1
    finally:
        wb.close()


# --------------------------------------------------------------------- 配对
def match_pairs(folder, sel_files, logger, tag, known=None):
    """
    返回 [(group, selected_path, viz_path), ...]，未匹配到的组写入日志。
    known: {selected 文件名: viz 文件名} 配对清单，优先复用
           （处理过后 selected 页已无原始腔室名，无法再按腔室集合匹配）。
    """
    src_folder = SRC_ROOT / folder.name
    cond = bev.en_labels.condition_of_folder(folder.name)
    if cond is None:
        logger.warning("%s 无法解析培养条件，跳过", tag)
        return []

    groups = []
    if src_folder.is_dir():
        for p in sorted(src_folder.glob("*.xlsx")):
            if p.name.startswith("~$") or not bev.XLSX_RE.search(p.name):
                continue
            try:
                groups.append(bev.read_group_xlsx(p, logger, tag))
            except Exception as exc:                            # noqa: BLE001
                logger.error("%s 读源失败 %s：%s", tag, p.name, exc)

    known = known or {}
    # selected → 原始腔室集合（已处理过的回退读源目录同名文件）
    sel_sets = {p: chamber_set_of(folder, p) for p in sel_files}

    pairs, used = [], set()
    for g in groups:
        src_set = g["src_set"]
        by_suffix = folder / (f"{cond['label']}_visualization_results"
                              f"{g['suffix']}.xlsx")
        picked, viz = None, by_suffix
        # ① 配对清单命中（重跑场景）：按 viz 文件名反查 selected
        hit = next((s for s in sel_files
                    if known.get(s.name) == by_suffix.name and s not in used),
                   None)
        if hit is not None:
            picked = hit
        # ② 按原始腔室集合匹配
        if picked is None and src_set:
            for p, names in sel_sets.items():
                if p in used:
                    continue
                if names == src_set:
                    picked = p
                    break
        # ③ 单组文件夹退化
        if picked is None and not src_set and len(groups) == 1 \
                and len(sel_files) == 1:
            picked = sel_files[0]
        if picked is None:
            logger.warning("%s %s 未匹配到 selected 文件（腔室数 %d）",
                           tag, Path(g["path"]).name, len(g["chambers"]))
            continue
        if not viz.exists():
            logger.warning("%s 找不到可视化工作簿 %s", tag, viz.name)
            continue
        used.add(picked)
        pairs.append((g, picked, viz))

    for p in sel_files:
        if p not in used:
            logger.warning("%s selected 文件 %s 未配对到任何数据组，跳过",
                           tag, p.name)
    return pairs


# --------------------------------------------------------------------- 主流程
def main(argv=None):
    parser = argparse.ArgumentParser(
        description="把可视化工作簿的培养条件区并入 selected__groups 文件")
    parser.add_argument("--dry-run", action="store_true", help="只打印计划")
    parser.add_argument("--force", action="store_true", help="重做已处理的文件")
    parser.add_argument("--only", default=None, help="只处理指定文件夹名")
    parser.add_argument("--no-backup", action="store_true", help="不建备份")
    args = parser.parse_args(argv)

    logger = setup_logger(config.PROJECT_ROOT / "simdata_pipeline" / "_logs",
                          name="selected_summary")
    logger.info("=" * 70)
    logger.info("selected__groups 并入培养条件区：%s", DST_ROOT)
    logger.info("模式：%s%s", "dry-run " if args.dry_run else "",
                "force" if args.force else "跳过已处理")

    if not DST_ROOT.is_dir():
        logger.error("目录不存在：%s", DST_ROOT)
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = BACKUP_ROOT / f"selected_cond_{stamp}"
    stats = {"folders": 0, "pairs": 0, "done": 0, "skipped": 0, "failed": 0}
    mapping = load_map()

    for folder in sorted(p for p in DST_ROOT.iterdir() if p.is_dir()):
        if args.only and folder.name != args.only:
            continue
        tag = f"[{folder.name}]"
        sel_files = sorted(p for p in folder.glob("selected__groups_*.xlsx")
                           if not p.name.startswith("~$"))
        if not sel_files:
            logger.info("%s 无 selected 文件，跳过", tag)
            continue
        stats["folders"] += 1

        known = {k.split("/", 1)[1]: v for k, v in mapping.items()
                 if k.startswith(folder.name + "/")}
        pairs = match_pairs(folder, sel_files, logger, tag, known=known)
        if not pairs:
            logger.warning("%s 无可处理配对", tag)
            continue

        for group, sel_path, viz_path in pairs:
            stats["pairs"] += 1
            ptag = f"{tag}{sel_path.name}"
            mapping[f"{folder.name}/{sel_path.name}"] = viz_path.name
            if not args.force and is_done(sel_path):
                logger.info("  %s 已处理过（单页签），跳过", ptag)
                stats["done"] += 1
                continue
            try:
                if not args.dry_run and not args.no_backup:
                    dst = backup_dir / folder.name / sel_path.name
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(sel_path, dst)
                res = selected_summary.apply_cond_block(
                    sel_path, viz_path, logger=logger, tag=ptag,
                    dry_run=args.dry_run)
                if args.dry_run:
                    logger.info("  %s [dry-run] 将复制 %s A1:H%d，"
                                "删除页签 %s，下移组合图",
                                ptag, viz_path.name, res["region_bottom"],
                                "/".join(res["old_sheets"]) or "无")
            except Exception as exc:                            # noqa: BLE001
                logger.error("  %s 处理失败：%s", ptag, exc)
                stats["failed"] += 1

    logger.info("=" * 70)
    logger.info("完成：文件夹 %d，配对 %d，处理 %d，已处理跳过 %d，失败 %d",
                stats["folders"], stats["pairs"],
                stats["pairs"] - stats["done"] - stats["failed"],
                stats["done"], stats["failed"])
    if not args.dry_run:
        if not args.only:
            save_map(mapping)
            logger.info("配对清单：%s（%d 条）", MAP_FILE, len(mapping))
        if not args.no_backup and stats["pairs"]:
            logger.info("备份：%s", backup_dir)
    return 1 if stats["failed"] else 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""
二次分组的产物生成
==================
复用两条既有链路，保证与「20组模拟数据 / 未使用数据」的产物完全同款：

    分组算法   ← batch_grouping/data_analysis_batch/grouping.allocate_groups
    分组工作簿 ← batch_grouping/data_analysis_batch/excel_io.build_group_workbooks
    PPT 版式   ← batch_grouping/data_analysis_batch/ppt_io.build_growth_curve_ppt
    可视化表格 ← simdata_pipeline/core/writer.build_workbook

唯一区别：PPT 用的腔室生长曲线 PNG 由本模块现场绘制（matplotlib），
因为 batch_grouping 的 extract_growth_charts 依赖源工作簿里已嵌入的 300 dpi
图片，而这些源文件的 drawing 关联用了绝对路径，openpyxl 读不出 _images。
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import config


# --------------------------------------------------------------------------- 复用 batch_grouping
def _batch_pkg():
    """导入 batch_grouping 的 data_analysis_batch 包（只读，不修改其文件）。"""
    root = Path(config.PROJECT_ROOT) / "batch_grouping"
    if not root.is_dir():
        raise FileNotFoundError(f"batch_grouping 目录不存在：{root}")
    key = str(root)
    if key not in sys.path:
        sys.path.insert(0, key)
    import data_analysis_batch as pkg
    return pkg


def allocate(rows, min_size=None, max_size=None, cv_limit=None, p_limit=None):
    """调用 batch_grouping 的分组算法，参数默认与 config 一致。"""
    pkg = _batch_pkg()
    from data_analysis_batch.grouping import allocate_groups, metrics
    from data_analysis_batch.models import ChamberResult

    results = []
    for r in rows:
        results.append(ChamberResult(
            source_file=r["source_file"],
            summary_row=r["summary_row"],
            chamber=r["chamber"],
            flow=r["flow"],
            flow_no=r["flow_no"],
            chamber_no=r["chamber_no"],
            umax=r["umax"], k=r["krel"], n0=r["n0"], r2=r["r2"],
            td=r["td"], f=r["f"], eta=r["eta"],
            lag=r["lag"], log=r["log"], stable=r["stable"],
        ))
    groups, remaining, full_count = allocate_groups(
        results,
        min_group_size=min_size or config.GROUP_MIN_SIZE,
        max_group_size=max_size or config.GROUP_MAX_SIZE,
        cv_limit=cv_limit if cv_limit is not None else config.GROUP_CV_LIMIT,
        normality_p_limit=p_limit if p_limit is not None else config.GROUP_P_LIMIT,
    )
    return results, groups, remaining, full_count, metrics


def group_stats(group, metrics):
    return metrics(group)


# --------------------------------------------------------------------------- 腔室曲线 PNG
def render_chamber_pngs(rows, out_dir, logger=None, tag=""):
    """
    为组内每个腔室绘制一张生长曲线 PNG（与可视化工作簿里的图同款）。

    返回 list[dict]：{order, label, image_path, source_file, worksheet}
    """
    import numpy as np

    from core import plotting
    from core.fitting import fit_curve

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = []
    for order, (row, ref) in enumerate(rows, start=1):
        counts = _numeric(ref.counts)
        t = np.arange(counts.size, dtype=float)
        if counts.size < 3:
            if logger:
                logger.warning("%s %s 数据点不足（%d），跳过出图",
                               tag, row.chamber, counts.size)
            continue
        try:
            fit = fit_curve(t, counts)
            buf = plotting.chamber_fit_figure(t, counts, fit, row.chamber)
        except Exception as exc:                                  # noqa: BLE001
            if logger:
                logger.warning("%s %s 出图失败：%s", tag, row.chamber, exc)
            continue
        path = out_dir / f"{order:02d}_{row.chamber}.png"
        path.write_bytes(buf.getvalue())
        charts.append({
            "order": order,
            "grid_row": (len(charts)) // 9,
            "grid_col": (len(charts)) % 9,
            "label": f"CH{row.flow_no}-CB{row.chamber_no}",
            "source_file": row.source_file,
            "worksheet": row.chamber,
            "image_path": str(path.resolve()),
        })
    # 网格坐标在跳过腔室后重新编号，保证 2 行 × 9 列连续排布
    for i, chart in enumerate(charts):
        chart["grid_row"] = i // 9
        chart["grid_col"] = i % 9
    return charts


def _numeric(values):
    import numpy as np
    return np.asarray([v for v in (values or []) if v is not None], dtype=float)


# --------------------------------------------------------------------------- 单组产物
def build_group_assets(group_index, rows, refs_by_key, folder, work_dir,
                       logger=None, tag="", plots=True, reference=None,
                       strict_rows=None):
    """
    rows        : 本组的 ChamberResult
    strict_rows : 全部分组候选的 ChamberResult（写入 strict_<N> 页签；不传则用 rows）
    """
    """
    为一个数据组生成全部产物。

        folder/
            <name>_可视化结果.xlsx
            selected__groups_<i>_<n>.xlsx
            <name>.pptx

    返回 dict（各产物文件名 + 统计）。
    """
    pkg = _batch_pkg()
    from data_analysis_batch.excel_io import build_group_workbooks, embed_chamber_images
    from data_analysis_batch.ppt_io import build_growth_curve_ppt

    from core.writer import build_workbook

    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    n = len(rows)
    result = {"size": n, "xlsx": None, "selected": None, "pptx": None, "notes": []}

    ordered = sorted(rows, key=lambda r: (r.flow_no, r.chamber_no, r.source_file))
    all_rows = list(strict_rows) if strict_rows else list(rows)

    # ---------- ① 可视化结果工作簿 ----------
    chambers = []
    for i, row in enumerate(ordered, start=1):
        ref = refs_by_key[(row.source_file, row.chamber)]
        chambers.append({
            "new_name": f"CH1_CB{i}",
            "rows": ref.rows,
            "counts": ref.counts,
            "areas": ref.areas,
            "params": params_of(ref),
            "src_file": ref.src_file,
            "src_sheet": ref.sheet,
            "src_rel": f"{ref.condition}/{ref.src_file}",
        })
    out_xlsx = folder / f"{folder.name}{config.GENERATED_SUFFIX}"
    try:
        build_workbook(out_xlsx, chambers, logger=logger, tag=tag,
                       condition=folder.name, plots=plots)
        result["xlsx"] = out_xlsx.name
    except Exception as exc:                                      # noqa: BLE001
        result["notes"].append(f"可视化工作簿失败：{exc}")
        if logger:
            logger.exception("%s 可视化工作簿生成失败：%s", tag, exc)

    # ---------- ② selected__groups_<i>_<n>.xlsx ----------
    ref_path = Path(reference or config.REFERENCE_XLSX)
    if not ref_path.is_file():
        msg = f"底稿不存在，跳过 selected__groups：{ref_path}"
        result["notes"].append(msg)
        if logger:
            logger.warning("%s %s", tag, msg)
    else:
        tmp_dir = Path(work_dir) / f"group_{group_index}"
        tmp_dir.mkdir(parents=True, exist_ok=True)
        source_paths = [Path(k) for k in sorted({r.source_file for r in all_rows})]
        try:
            build_group_workbooks(
                ref_path, tmp_dir, [ordered], all_rows, source_paths,
                config.GROUP_MIN_SIZE, config.GROUP_MAX_SIZE,
            )
            produced = sorted(tmp_dir.glob("selected__groups_1_*.xlsx"))
            if produced:
                target = folder / f"selected__groups_{group_index}_{n}.xlsx"
                shutil.move(str(produced[0]), str(target))
                _patch_source_files(target, [p.as_posix() for p in source_paths])
                result["selected"] = target.name
        except Exception as exc:                                  # noqa: BLE001
            result["notes"].append(f"分组工作簿失败：{exc}")
            if logger:
                logger.exception("%s 分组工作簿生成失败：%s", tag, exc)

    # ---------- ③ PPT ----------
    charts = render_chamber_pngs(
        [(row, refs_by_key[(row.source_file, row.chamber)]) for row in ordered],
        Path(work_dir) / "ppt_charts" / f"group_{group_index}",
        logger=logger, tag=tag)
    if not charts:
        result["notes"].append("无可用腔室曲线，未生成 pptx")
    else:
        try:
            stats = pkg.grouping.metrics(ordered)
            manifest = {"groups": [{
                "index": group_index,
                "size": n,
                "source_workbook": f"selected__groups_{group_index}_{n}.xlsx",
                "selected_sheet": f"selected_{n}",
                "umax_mean_h-1": stats.mean,
                "umax_cv_percent": stats.cv * 100,
                "charts": charts,
            }]}
            pptx_path = folder / config.SECOND_PPTX_NAME_TEMPLATE.format(folder=folder.name)
            build_growth_curve_ppt(pptx_path, manifest)
            result["pptx"] = pptx_path.name
        except Exception as exc:                                  # noqa: BLE001
            result["notes"].append(f"PPT 生成失败：{exc}")
            if logger:
                logger.exception("%s PPT 生成失败：%s", tag, exc)

        # 复合图嵌入分组工作簿（与 batch_grouping 的 embed_chamber_images 一致）
        if result["selected"]:
            try:
                embed_chamber_images(folder / result["selected"], charts)
            except Exception as exc:                              # noqa: BLE001
                result["notes"].append(f"复合图嵌入失败：{exc}")
                if logger:
                    logger.warning("%s 复合图嵌入失败：%s", tag, exc)

    return result


def _patch_source_files(path, names):
    """build_group_workbooks 只写文件名，这里补全「工况/文件」的完整来源。"""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb["summary"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "source_files":
            ws.cell(r, 2).value = "; ".join(names)
            break
    wb.save(path)
    wb.close()


def params_of(ref):
    """腔室页签的「参数名 / 参数值」两列 → dict。"""
    params = {}
    for r in ref.rows or []:
        if len(r) >= 6 and r[4] is not None:
            key = str(r[4]).strip()
            if key and key != "参数名":
                params[key] = r[5]
    return params


SHORT_FILE_RE = re.compile(r"_可视化结果(_\d+)?$")


def short_file(name: str) -> str:
    """20260504_CH1_可视化结果_96.xlsx → 20260504_CH1"""
    stem = name[:-5] if name.lower().endswith(".xlsx") else name
    return SHORT_FILE_RE.sub("", stem)


def short_condition(cond: str) -> str:
    """L120‑N160‑IC0.5% → L120-N160-IC0.5（去掉 % 与 U+2011，便于做页签名）"""
    return cond.replace("\u2011", "-").replace("\u2010", "-").replace("%", "")

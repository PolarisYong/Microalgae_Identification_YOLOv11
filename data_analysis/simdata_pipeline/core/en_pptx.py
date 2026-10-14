# -*- coding: utf-8 -*-
"""
英文版分组 PPT 重建
===================
以源 PPT 为底稿（版式 / 位置 / 尺寸不动），做三件事：

1. 页眉文字英文化：「数据组 1（n = 18）」→「Group 1 (n = 18)」（其余不动）
2. 腔室标签改写：CH4-CB1 / CB1 → CB-<k>（k 为该腔室在新工作簿中的序号）
3. 配图原地重绘：Y = ln(Nrel) 尺度、英文标题 Growth curve (CB-<k>, <cond>)，
   按原图像素尺寸渲染后替换图片部件二进制

腔室与图片的对应关系：优先用「图片正上方同列标签框」里的原标签
（CH<x>-CB<y> 经数据来源映射，或 CB<n> 直接取序号）定位腔室；
没有标签时退化为按 (top, left) 几何顺序对应腔室顺序。
"""
import re
from pathlib import Path

from core import en_labels, en_plotting

# PPT 标签：CH11-CB36 或 CB12
LABEL_SOURCE_RE = re.compile(r"^CH(\d+)-CB(\d+)$")
LABEL_CB_RE = re.compile(r"^CB-?(\d+)$")
MEAN_LINE_RE = re.compile(r"^Mean μmax = ([\d.]+) h⁻¹$")
GROUP_LINE_RE = re.compile(r"^Group (\d+) \(n = (\d+)\)$")


def _rich_symbol_runs(shape):
    """
    把含符号的文本框改写成斜体 + 上下标的多 run 文本：

    - 「Mean μmax = 0.0658 h⁻¹」→ Mean + μ(斜) + max(下标) + = 0.0658 h + -1(上标)
    - 「Group 1 (n = 18)」      → n 斜体

    保留原 run 的字号 / 粗体；字体由后续统一步骤设为 Times New Roman。
    """
    text = shape.text_frame.text.strip()
    m = MEAN_LINE_RE.match(text)
    if m:
        parts = [("Mean ", {}), ("μ", {"i": 1}), ("max", {"b": "-25000"}),
                 (f" = {m.group(1)} h", {}), ("-1", {"b": "30000"})]
    else:
        m = GROUP_LINE_RE.match(text)
        if not m:
            return False
        parts = [(f"Group {m.group(1)} (", {}), ("n", {"i": 1}),
                 (f" = {m.group(2)})", {})]

    para = shape.text_frame.paragraphs[0]
    ref = para.runs[0] if para.runs else None
    size = ref.font.size if ref is not None else None
    bold = ref.font.bold if ref is not None else None
    para.clear()
    for txt, style in parts:
        run = para.add_run()
        run.text = txt
        if size is not None:
            run.font.size = size
        if bold is not None:
            run.font.bold = bold
        if style.get("i"):
            run.font.italic = True
        if style.get("b"):
            run._r.get_or_add_rPr().set("baseline", style["b"])  # noqa: SLF001
    return True


def _shape_key(shape):
    return (shape.top, shape.left)


def _set_text(shape, text):
    """保留首个 run 的字体格式，替换整段文字。"""
    tf = shape.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        para.text = text


def _label_of(shape):
    """标签框文本 → ('src', 'CH11_CB36') / ('cb', 12) / None。"""
    text = shape.text_frame.text.strip()
    m = LABEL_SOURCE_RE.fullmatch(text)
    if m:
        return ("src", f"CH{m.group(1)}_CB{m.group(2)}")
    m = LABEL_CB_RE.fullmatch(text)
    if m:
        return ("cb", int(m.group(1)))
    return None


def _as_queue(mapping):
    """
    {原始腔室名: 序号} 或 {原始腔室名: [序号, ...]} → {名: deque([序号...])}。
    同一原始腔室被重复纳入一组时（如 CH1_CB23 ×2），按出现顺序依次消费。
    """
    from collections import deque

    out = {}
    for key, value in (mapping or {}).items():
        if isinstance(value, (list, tuple)):
            out[key] = deque(value)
        else:
            out[key] = deque([value])
    return out


def rebuild_pptx_en(src_pptx, dst_pptx, chambers, cond, src_to_index,
                    logger=None, tag=""):
    """
    chambers     : 新工作簿腔室列表（含 index / cb_id / counts / params）
    cond         : en_labels.condition_of_folder 的返回
    src_to_index : {原始腔室名: 序号} 或 {原始腔室名: [序号...]}（重复腔室按序消费）
    返回 (替换图片数, 改写标签数, 翻译页眉数, 统一字体 run 数)。
    """
    import numpy as np
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(src_pptx))
    if not prs.slides:
        raise ValueError(f"{Path(src_pptx).name} 里没有幻灯片")
    slide = prs.slides[0]

    pictures = sorted([sh for sh in slide.shapes
                       if sh.shape_type == MSO_SHAPE_TYPE.PICTURE], key=_shape_key)
    label_shapes = [sh for sh in slide.shapes
                    if sh.has_text_frame and _label_of(sh) is not None]
    if len(pictures) != len(chambers):
        raise ValueError(
            f"PPT 里有 {len(pictures)} 张图，但工作簿有 {len(chambers)} 个腔室，"
            f"数量不一致，已中止以免错配")

    by_index = {ch["index"]: ch for ch in chambers}
    src_queue = _as_queue(src_to_index)

    def pair_label(pic):
        """
        图片正上方的标签框。源 PPT 里标签与图片的 left 存在 1~2 EMU 的
        微小偏差，严格相等会漏配；改为「标签水平中心落在图片水平区间内、
        垂直方向位于图片上方、距离最近」。
        """
        target, best = None, None
        for lab in label_shapes:
            center_x = lab.left + (lab.width or 0) / 2
            if not (pic.left - 20000 <= center_x <= pic.left + pic.width + 20000):
                continue
            distance = pic.top - (lab.top + (lab.height or 0))
            if distance < -20000:                # 标签明显压进图片区域 → 不是它的标签
                continue
            if best is None or distance < best:
                target, best = lab, distance
        return target

    # ① 逐张图：确定腔室 → 重绘英文 Y 尺度图 → 替换二进制
    relabeled = 0
    for order, pic in enumerate(pictures, 1):
        lab = pair_label(pic)
        info = _label_of(lab) if lab is not None else None
        index = None
        if info and info[0] == "src":
            queue = src_queue.get(info[1])
            index = queue.popleft() if queue else None
            if index is None and logger:
                logger.warning("%s 标签 %s 未在数据来源中匹配到腔室，"
                               "退化为按顺序第 %d 个", tag, info[1], order)
        elif info and info[0] == "cb":
            index = info[1] if info[1] in by_index else None
        if index is None:
            index = order
            if logger:
                logger.warning("%s 第 %d 张图未配对到标签，按顺序取腔室 %d",
                               tag, order, index)

        ch = by_index.get(index)
        if ch is None:
            raise ValueError(f"第 {order} 张图对应腔室序号 {index} 不存在")

        counts = np.asarray([v for v in ch["counts"] if v is not None],
                            dtype=float)
        t = np.arange(counts.size, dtype=float)
        size_px = pic.image.size
        png = en_plotting.ppt_chamber_figure_en(
            t, counts, ch["params"], ch["cb_id"], cond["title_label"], size_px)

        r_id = pic._element.blipFill.blip.rEmbed          # noqa: SLF001
        part = pic.part.related_part(r_id)
        part._blob = png.read()                           # noqa: SLF001
        if hasattr(part, "_image"):
            del part._image

        if lab is not None:
            _set_text(lab, ch["cb_id"])
            relabeled += 1

    # ② 页眉文字英文化 + 符号富文本（斜体/上下标）
    headers = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text
        new_text = en_labels.translate_ppt_header(text)
        if new_text != text:
            _set_text(sh, new_text)
            headers += 1
        if _rich_symbol_runs(sh):
            headers += 1

    # ③ 字体统一：全部文本 run → Times New Roman（原为 Arial / 微软雅黑）
    unified = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.text and run.font.name != "Times New Roman":
                    run.font.name = "Times New Roman"
                    unified += 1

    prs.save(str(dst_pptx))
    return len(pictures), relabeled, headers, unified

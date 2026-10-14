# -*- coding: utf-8 -*-
"""
扫描「尚未参与数据分组」的腔室
==============================
遍历 `数据汇总分析结果/<自变量文件夹>/` 下的原始 .xlsx，读取其「汇总结果」页签；
再用同目录 `18_batch_*/<同名>_marked.xlsx` 的**标红行**判定哪些腔室已被使用
（红色字体 = 已参与 batch_grouping 输出的数据组）。未被标红的腔室即为本轮的输入。

只读，不修改任何源文件。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import config

CHAMBER_RE = re.compile(r"^CH(\d+)_CB(\d+)$")
SUMMARY_SHEET = "汇总结果"
RED_TAIL = "FF0000"

# 汇总结果列索引（0-based，与 batch_grouping/data_analysis_batch/excel_io.collect_candidates 一致）
COL_CHAMBER = 0
COL_umax, COL_KREL, COL_N0, COL_R2 = 1, 2, 3, 4
COL_TD, COL_F, COL_ETA, COL_LAG, COL_LOG, COL_STABLE = 5, 6, 7, 8, 9, 10

# 腔室页签列索引
COL_IMAGE, COL_STATUS, COL_COUNT, COL_AREA = 0, 1, 2, 3


@dataclass
class ChamberRef:
    """一个未参与分组的腔室，含汇总指标 + 原始页签内容。"""

    condition: str          # 自变量文件夹名
    src_file: str           # 源 .xlsx 文件名
    sheet: str              # 原始页签名 CH<x>_CB<y>
    flow_no: int
    chamber_no: int
    umax: float
    krel: float
    n0: float
    r2: float
    td: float
    f: float
    eta: float
    lag: float
    log: float
    stable: float
    summary_row: int
    rows: list = field(default_factory=list)      # 腔室页签全部单元格值
    counts: list = field(default_factory=list)    # 「目标数量」列
    areas: list = field(default_factory=list)     # 「总面积(μm²)」列
    fail_count: int = 0                           # 「处理状态」非成功的行数

    @property
    def uid(self) -> tuple:
        return (self.condition, self.src_file, self.sheet)

    @property
    def origin(self) -> str:
        return f"{self.condition}｜{self.src_file}｜{self.sheet}"


def _num(value):
    if value is None or value == "":
        return None
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    return v if v == v and abs(v) != float("inf") else None


def is_source_workbook(path: Path) -> bool:
    """是否是参与分组的原始工作簿（排除 marked / selected / 临时文件）。"""
    name = path.name
    if not name.lower().endswith(".xlsx"):
        return False
    if name.startswith("~$") or name.startswith("selected__groups_"):
        return False
    if name.endswith("_marked.xlsx") or name == "参考输出.xlsx":
        return False
    return True


def find_marked(cond_dir: Path, src_name: str, batch_prefix: str = None,
                batch_name: str = None):
    """定位源文件的 _marked.xlsx；找不到返回 None。"""
    prefix = batch_prefix or config.BATCH_DIR_PREFIX
    stem = src_name[:-5]            # 去掉 .xlsx
    candidates = []
    if batch_name:
        candidates.append(cond_dir / batch_name)
    for d in sorted(p for p in cond_dir.iterdir()
                    if p.is_dir() and p.name.startswith(prefix)):
        if d not in candidates:
            candidates.append(d)
    for d in candidates:
        p = d / f"{stem}_marked.xlsx"
        if p.is_file():
            return p
    return None


def _red_chambers(marked: Path) -> set:
    """读取 marked 工作簿「汇总结果」页签里字体为红色的腔室名。"""
    import openpyxl

    wb = openpyxl.load_workbook(marked, read_only=True, data_only=False)
    try:
        if SUMMARY_SHEET not in wb.sheetnames:
            return set()
        ws = wb[SUMMARY_SHEET]
        red = set()
        for row in ws.iter_rows(min_row=2, values_only=False):
            cell = row[COL_CHAMBER]
            name = cell.value
            if not isinstance(name, str) or not CHAMBER_RE.fullmatch(name):
                continue
            color = cell.font.color
            rgb = getattr(color, "rgb", None) if color is not None else None
            if rgb is not None and str(rgb).upper().endswith(RED_TAIL):
                red.add(name)
        return red
    finally:
        wb.close()


def _read_chamber_sheet(ws) -> tuple:
    """腔室页签 → (rows, counts, areas, fail_count)。"""
    rows, counts, areas, fail = [], [], [], 0
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
        status = r[COL_STATUS] if len(r) > COL_STATUS else None
        if status is not None and str(status).strip() != "成功":
            fail += 1
        v = _num(r[COL_COUNT]) if len(r) > COL_COUNT else None
        a = _num(r[COL_AREA]) if len(r) > COL_AREA else None
        counts.append(v)
        areas.append(a)
    return rows, counts, areas, fail


def scan_unused(agg_dir, logger=None, batch_name: str = None) -> dict:
    """
    扫描全部自变量文件夹。

    返回 dict：
        refs        : list[ChamberRef]  未参与分组的腔室
        used        : int               被标红（已使用）的腔室数
        files       : int               读取的源工作簿数
        conditions  : int               自变量文件夹数
        no_marked   : list[str]         找不到 _marked.xlsx 的源文件名（全部腔室按未使用处理）
        missing_sheet: list[str]        汇总结果里有、但工作簿中缺失的腔室页签
    """
    import openpyxl

    agg_dir = Path(agg_dir)
    if not agg_dir.is_dir():
        raise FileNotFoundError(f"汇总结果目录不存在：{agg_dir}")

    batch_name = batch_name or getattr(config, "BATCH_DIR_NAME", "18_batch_20260915")
    refs, no_marked, missing_sheet = [], [], []
    used = files = conditions = 0

    for cond_dir in sorted(p for p in agg_dir.iterdir() if p.is_dir()):
        sources = [p for p in sorted(cond_dir.iterdir())
                   if p.is_file() and is_source_workbook(p)]
        if not sources:
            continue
        conditions += 1

        for src in sources:
            marked = find_marked(cond_dir, src.name, batch_name=batch_name)
            red = set()
            if marked is None:
                no_marked.append(f"{cond_dir.name}/{src.name}")
                if logger:
                    logger.warning("[%s] 找不到 _marked.xlsx，该文件全部腔室按「未使用」处理：%s",
                                   cond_dir.name, src.name)
            else:
                red = _red_chambers(marked)

            wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
            try:
                if SUMMARY_SHEET not in wb.sheetnames:
                    if logger:
                        logger.warning("[%s] 缺少「%s」页签，已跳过：%s",
                                       cond_dir.name, SUMMARY_SHEET, src.name)
                    continue
                ws = wb[SUMMARY_SHEET]
                for r_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True),
                                               start=2):
                    name = values[COL_CHAMBER] if values else None
                    if not isinstance(name, str):
                        continue
                    m = CHAMBER_RE.fullmatch(name)
                    if not m:
                        continue
                    if name in red:
                        used += 1
                        continue
                    if name not in wb.sheetnames:
                        missing_sheet.append(f"{cond_dir.name}/{src.name}｜{name}")
                        continue

                    nums = [values[i] if i < len(values) else None
                            for i in range(COL_umax, COL_STABLE + 1)]
                    nums = [_num(v) for v in nums]
                    rows, counts, areas, fail = _read_chamber_sheet(wb[name])
                    # 表头行不算数据点
                    data_counts = [v for v in counts[1:] if v is not None]
                    data_areas = [v for v in areas[1:] if v is not None]

                    refs.append(ChamberRef(
                        condition=cond_dir.name,
                        src_file=src.name,
                        sheet=name,
                        flow_no=int(m.group(1)),
                        chamber_no=int(m.group(2)),
                        umax=nums[0] if nums[0] is not None else float("nan"),
                        krel=nums[1] if nums[1] is not None else float("nan"),
                        n0=nums[2] if nums[2] is not None else float("nan"),
                        r2=nums[3] if nums[3] is not None else float("nan"),
                        td=nums[4] if nums[4] is not None else float("nan"),
                        f=nums[5] if nums[5] is not None else float("nan"),
                        eta=nums[6] if nums[6] is not None else float("nan"),
                        lag=nums[7] if nums[7] is not None else float("nan"),
                        log=nums[8] if nums[8] is not None else float("nan"),
                        stable=nums[9] if nums[9] is not None else float("nan"),
                        summary_row=r_idx,
                        rows=rows,
                        counts=data_counts,
                        areas=data_areas,
                        fail_count=fail,
                    ))
                files += 1
            finally:
                wb.close()

    if logger:
        logger.info("扫描完成：%d 个自变量文件夹 / %d 个源工作簿｜已使用（标红）%d 个腔室"
                    "｜未参与分组 %d 个腔室",
                    conditions, files, used, len(refs))
        if missing_sheet:
            logger.warning("汇总结果里有记录但工作簿中缺页签的腔室：%d 个（已跳过）",
                           len(missing_sheet))
            for s in missing_sheet[:10]:
                logger.warning("  · %s", s)
    return {
        "refs": refs,
        "used": used,
        "files": files,
        "conditions": conditions,
        "no_marked": no_marked,
        "missing_sheet": missing_sheet,
    }

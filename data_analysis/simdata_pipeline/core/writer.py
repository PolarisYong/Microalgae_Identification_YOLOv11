# -*- coding: utf-8 -*-
"""
生成输出工作簿
==============
每个工况（自变量组合）一个 .xlsx，内含三类页签：

    汇总        —— 三张表，布局对齐参考实现 08_processed_batch.py 的「汇总结果」：
                  ① 腔室级拟合汇总：每个腔室一行，字段见 config.SUMMARY_FIELDS
                  ② 未参与拟合数据：仅当存在数据不完整的腔室时写入（红色分隔行）
                  ③ 合并拟合结果：全部可用腔室拼接后一次性拟合
                     （**不做腔室筛选**：没有后向剔除 / beam search / 回看修复）
                  与「合并拟合结果」同一行、向右横向排开另嵌 3 张 matplotlib 图
                  （生长曲线 / 相对总面积 / 相对平均面积），列间距按列宽自动计算，互不重叠
    CH1_CB1..N  —— 原始数据页签，腔室统一重命名为 CH1_CB<序号>；
                  每页右侧嵌 4 张图（单腔室拟合 / 与合并曲线对比 / 总面积 / 平均面积）
    数据来源    —— 【腔室名】【原始数据来源】两列，记录取自哪个 excel 的哪个页签

合并拟合的口径与参考实现一致：模型拟合在 Y=ln(Nrel) 尺度，R² 同尺度计算。
"""
from pathlib import Path

import numpy as np

import config
from core.fitting import merged_fit
from core.rawdata import relative_average_area


def _summary_row(ch):
    """腔室 → 汇总表一行的取值列表。"""
    params = ch.get("params") or {}
    return [ch["new_name"] if src_key is None else params.get(src_key)
            for _h, src_key in config.SUMMARY_FIELDS]


def _merged_values(merge):
    """合并拟合结果 → 与 config.SUMMARY_MERGED_FIELDS 一一对应的取值列表。"""
    fit = merge["fit"]
    phase = merge["phase_duration"]
    t90 = merge["t90"]
    return [
        len(merge["valid_sheets"]),
        int(len(merge["merged_t"])),
        round(fit["mu_max"], 4),
        round(fit["K"], 2),
        round(fit["N0"], 2),
        round(fit["lambda_"], 2),
        round(fit["r2"], 4),
        phase["滞后期时长(h)"],
        phase["对数期时长(h)"],
        phase["稳定期时长(h)"],
        round(float(t90), 2) if np.isfinite(t90) else None,
    ]


def _write_summary(wb, chambers, merge):
    """写出「汇总」页签的三张表。"""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("汇总")
    headers = [h for h, _ in config.SUMMARY_FIELDS]
    n_col = len(headers)

    # ---------- 表① 腔室级拟合汇总 ----------
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ch in chambers:
        ws.append(_summary_row(ch))

    # ---------- 表② 未参与拟合数据（仅在存在时写入） ----------
    excluded = set(merge["excluded"])
    if excluded:
        r = ws.max_row + 2
        ws.cell(row=r, column=1, value="未参与拟合数据")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_col)
        ws.cell(row=r, column=1).font = Font(bold=True, color="FF0000")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        for ch in chambers:
            if ch["new_name"] in excluded:
                ws.append(_summary_row(ch))

    # ---------- 表③ 合并拟合结果 ----------
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="合并拟合结果").font = Font(bold=True)
    merged_row = r                      # 汇总页 3 张图的顶边对齐这一行
    if merge["fit"] is None:
        ws.cell(row=r + 1, column=1,
                value="（未能拟合：可用腔室不足 2 个，或数据含非正值）")
    else:
        for i, (label, value) in enumerate(
                zip(config.SUMMARY_MERGED_FIELDS, _merged_values(merge))):
            ws.cell(row=r + 1 + i, column=1, value=label)
            ws.cell(row=r + 1 + i, column=2, value=value)

    # ---------- 列宽 ----------
    ws.column_dimensions["A"].width = 26
    for i in range(2, n_col + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
    ws.column_dimensions["D"].width = 24
    return ws, merged_row


def _numeric(values):
    """剔除空值并转 float。"""
    try:
        return np.asarray([v for v in (values or []) if v is not None], dtype=float)
    except (TypeError, ValueError):
        return np.array([], dtype=float)


def _chamber_series(ch):
    """腔室 dict → (counts, areas)，均为一维 float 数组。"""
    return _numeric(ch.get("counts")), _numeric(ch.get("areas"))


def _add_chamber_plots(ws, ch, merge, plt_name, logger, tag):
    """腔室页右侧嵌入 4 张图（锚点 H2 / H24 / R2 / R24）。"""
    from core import plotting
    from core.fitting import fit_curve

    counts, areas = _chamber_series(ch)
    if counts.size < 3:
        if logger:
            logger.warning("%s %s 数据点不足（%d），跳过绘图",
                           tag, ch["new_name"], counts.size)
        return
    t = np.arange(counts.size, dtype=float)
    anchors = config.PLOT_CHAMBER_ANCHORS

    # 图 1：单腔室生长曲线拟合（红）
    try:
        fit = fit_curve(t, counts)
        plotting.insert_image(ws, plotting.chamber_fit_figure(t, counts, fit, plt_name),
                              anchors[0], config.PLOT_CHAMBER_SIZE)
    except Exception as exc:                                  # noqa: BLE001
        if logger:
            logger.warning("%s %s 单腔室拟合图生成失败：%s", tag, ch["new_name"], exc)

    # 图 2：该腔室数据 vs 合并拟合曲线（黄）
    if merge.get("fit") is not None:
        try:
            plotting.insert_image(
                ws, plotting.chamber_vs_merged_figure(t, counts, merge["fit"], plt_name),
                anchors[1], config.PLOT_CHAMBER_SIZE)
        except Exception as exc:                              # noqa: BLE001
            if logger:
                logger.warning("%s %s 合并对比图生成失败：%s", tag, ch["new_name"], exc)

    # 图 3/4：相对总面积（绿）、相对平均面积（紫）
    n = min(config.PLOT_AREA_POINTS, areas.size, counts.size)
    if n >= 3:
        try:
            total_buf, avg_buf = plotting.chamber_area_figures(
                t[:n], areas[:n], relative_average_area(areas[:n], counts[:n]), plt_name)
            plotting.insert_image(ws, total_buf, anchors[2], config.PLOT_CHAMBER_SIZE)
            plotting.insert_image(ws, avg_buf, anchors[3], config.PLOT_CHAMBER_SIZE)
        except Exception as exc:                              # noqa: BLE001
            if logger:
                logger.warning("%s %s 面积图生成失败：%s", tag, ch["new_name"], exc)


# 表格自身占用的列，排图时不能改它们的宽度（会挤压表头）
_PRESERVE_WIDTH_COLUMNS = ("A", "D")


def _column_px(ws, index):
    """列宽（openpyxl 的"字符"单位）→ 像素。Excel 换算：px ≈ 字符宽 * 7 + 5。"""
    from openpyxl.utils import get_column_letter

    dim = ws.column_dimensions.get(get_column_letter(index))
    width = getattr(dim, "width", None) or 8.43      # 未设宽度的列按默认 8.43 算
    return float(width) * 7 + 5


def _merged_anchors(ws, row, start_index, count):
    """
    从 start_index 列起，为 count 张图算出互不重叠的锚点。

    每张图占 config.PLOT_MERGED_SIZE[0] 像素，之后向右跳过若干整列，
    直到累计宽度 ≥ 图宽 + config.PLOT_MERGED_GAP_PX，下一张图从该列开始。
    """
    from openpyxl.utils import get_column_letter

    img_w = config.PLOT_MERGED_SIZE[0]
    need = img_w + config.PLOT_MERGED_GAP_PX
    anchors, idx = [], start_index
    for _ in range(count):
        anchors.append(f"{get_column_letter(idx)}{row}")
        accumulated = 0.0
        while accumulated < need:
            accumulated += _column_px(ws, idx)
            idx += 1
    return anchors


def _prepare_plot_band(ws, start_index, end_index):
    """
    统一图区域的列宽。未显式设宽度的列在 Excel 里只有 8.43 字符（约 64 px），
    三张 700 px 的图会挤在一起，这里补齐成 config.PLOT_MERGED_BAND_WIDTH。
    """
    from openpyxl.utils import get_column_letter

    for i in range(start_index, end_index + 1):
        letter = get_column_letter(i)
        if letter in _PRESERVE_WIDTH_COLUMNS:
            continue
        ws.column_dimensions[letter].width = config.PLOT_MERGED_BAND_WIDTH


def _add_merged_plots(ws, chambers, merge, condition, logger, tag, row):
    """
    汇总页嵌入 3 张图：顶边对齐「合并拟合结果」所在行，向右依次排开。
    row 由 _write_summary 返回（即「合并拟合结果」标题行）。
    """
    from openpyxl.utils import column_index_from_string

    from core import plotting

    fit = merge.get("fit")
    if fit is None:
        return
    valid = set(merge["valid_sheets"])
    series = []
    for ch in chambers:
        if ch["new_name"] not in valid:
            continue
        counts, areas = _chamber_series(ch)
        if counts.size < 3:
            continue
        series.append((counts, areas))
    if not series:
        return

    n_area = config.PLOT_AREA_POINTS
    merged_t = np.concatenate([np.arange(c.size, dtype=float) for c, _ in series])
    merged_counts = np.concatenate([c for c, _ in series])
    t_area = np.concatenate([np.arange(min(c.size, n_area), dtype=float) for c, _ in series])
    area_all = np.concatenate([a[:n_area] for _, a in series])
    avg_all = np.concatenate([relative_average_area(a[:n_area], c[:n_area])
                              for c, a in series])
    n_chambers = len(merge["valid_sheets"])

    row = row + config.PLOT_MERGED_ROW_OFFSET
    start = column_index_from_string(config.PLOT_MERGED_START_COLUMN)
    _prepare_plot_band(ws, start, start + 24)       # 预留足够列，避免宽度不一致导致重叠
    cols = _merged_anchors(ws, row, start, 3)
    try:
        plotting.insert_image(
            ws, plotting.merged_growth_figure(merged_t, merged_counts, fit,
                                              n_chambers, condition),
            cols[0], config.PLOT_MERGED_SIZE)
        total_buf, avg_buf = plotting.merged_area_figures(
            t_area, area_all, avg_all, n_chambers, condition)
        plotting.insert_image(ws, total_buf, cols[1], config.PLOT_MERGED_SIZE)
        plotting.insert_image(ws, avg_buf, cols[2], config.PLOT_MERGED_SIZE)
    except Exception as exc:                                  # noqa: BLE001
        if logger:
            logger.warning("%s 汇总页图生成失败：%s", tag, exc)


def build_workbook(out_path, chambers, logger=None, tag="", condition="", plots=None,
                   include_source=True):
    """
    chambers: list of dict
        {'new_name': 'CH1_CB1',
         'rows': [[...]],                      # 原始页签内容
         'counts': [1.0, 1.03, ...],           # 「目标数量」列，供合并拟合使用
         'areas':  [1.0, 1.02, ...],           # 「总面积(μm²)」列（已按 0 h 归一化）
         'params': {参数名: 参数值},
         'src_file': '20260712_CH7_可视化结果_96.xlsx',
         'src_sheet': 'CH7_CB7',
         'src_rel':  '20260712\\数据汇总\\...'}

    condition : 工况名（文件夹名），用于图标题
    plots     : None 时取 config.PLOT_ENABLED
    include_source : False 则不写「数据来源」页签（用于需要隐去原始出处的改写场景）

    返回 (写入路径, 实际腔室数, 合并拟合结果)
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font

    plots = config.PLOT_ENABLED if plots is None else plots
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    # 合并拟合：全部腔室直接拼接，不做筛选
    merge = merged_fit([(ch["new_name"], ch.get("counts") or []) for ch in chambers])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_summary, merged_row = _write_summary(wb, chambers, merge)

    # ---------------- 腔室页签（数据 + 4 张图） ----------------
    for ch in chambers:
        ws = wb.create_sheet(ch["new_name"])
        for r in ch.get("rows") or []:
            ws.append(list(r))
        if plots:
            # 图标题里的腔室标签：优先取腔室字典里的 plot_name（改写场景直接给全名），
            # 否则沿用「工况_ + 页签名去掉 CH<k>_ 前缀」的旧规则。
            plt_name = ch.get("plot_name") or (
                f"{condition}_{ch['new_name'][4:]}" if condition else ch["new_name"])
            _add_chamber_plots(ws, ch, merge, plt_name, logger, tag)

    # ---------------- 数据来源 ----------------
    if include_source:
        ws_src = wb.create_sheet("数据来源")
        ws_src.append(["腔室名", "原始数据来源"])
        for c in ws_src[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for ch in chambers:
            ws_src.append([ch["new_name"],
                           f"{ch.get('src_rel', ch.get('src_file', ''))}｜页签：{ch['src_sheet']}"])

        ws_src.column_dimensions["A"].width = 14
        ws_src.column_dimensions["B"].width = 90

    # ---------------- 汇总页的 3 张图（须在表格写完后定位行号） ----------------
    if plots:
        _add_merged_plots(ws_summary, chambers, merge, condition, logger, tag,
                          row=merged_row)

    wb.save(out_path)

    if logger:
        if merge["fit"] is None:
            logger.warning("%s 合并拟合未成功（可用腔室 %d 个）→ %s",
                           tag, len(merge["valid_sheets"]), out_path.name)
        else:
            logger.info("  %s 写出 %s（%d 个腔室页签%s；合并拟合：%d 个腔室 / %d 个数据点，"
                        "μmax=%.4f h⁻¹，R²=%.4f）",
                        tag, out_path.name, len(chambers),
                        "，共 %d 张图" % (len(chambers) * 4 + 3) if plots else "",
                        len(merge["valid_sheets"]), len(merge["merged_t"]),
                        merge["fit"]["mu_max"], merge["fit"]["r2"])
        if merge["excluded"]:
            logger.warning("  %s 未参与拟合的腔室：%s", tag, ", ".join(merge["excluded"]))

    return out_path, len(chambers), merge

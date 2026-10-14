# -*- coding: utf-8 -*-
"""
英文版工作簿生成
================
按「可视化结果参考模板.xlsx」生成英文版工作簿（不重新拟合，全部使用
源工作簿里已存的拟合参数）：

    汇总结果    —— ① 8 列腔室表（CB-1..，Tlag = λ、Tprolif = t90 − λ 新口径）
                  ② 培养条件块（coded 三元组 + L/N/DIC 实际值）
                  ③ Pooled fitting results 11 行
                  右侧嵌 3 张合并图（Y 尺度生长 / Atot,rel / Amean,rel，
                  锚点按列宽累加防重叠，复用 core.writer 的几何逻辑）
    CB-1..CB-n  —— A:F 数据表 + E/F 10 行参数块 + 4 张图
                  （A 列图片名 CB{k}_H{i}.tif，状态 成功→Success）
"""
from pathlib import Path

import numpy as np

import config
from core import en_labels, en_plotting, en_richtext
from core.rawdata import relative_average_area
from core.writer import _merged_anchors, _prepare_plot_band
from core.plotting import insert_image


# --------------------------------------------------------------------- 参数换算
def derived_params(p, t_end):
    """
    源参数 {mu,K,N0,lam,r2,td,tplat,t90} → 模板 10 行的取值列表。
    新口径（用户确认）：Tlag = λ（可为负），Tprolif = t90 − λ。
    """
    lam = _f(p.get("lam"))
    t90 = _f(p.get("t90"))
    tlag = lam
    tprolif = (t90 - lam) if (lam is not None and t90 is not None) else None
    tplat = p.get("tplat")
    if tplat is None and t90 is not None and t_end is not None:
        tplat = t_end - t90
    return [
        _r(p.get("mu"), 4),
        _r(p.get("K"), 4),
        _r(p.get("N0"), 4),
        _r(lam, 4),
        _r(p.get("r2"), 4),
        _r(p.get("td"), 2),
        _r(tlag, 2),
        _r(tprolif, 2),
        _r(tplat, 2),
        _r(t90, 2),
    ]


def _f(value):
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    return v if np.isfinite(v) else None


def _r(value, ndigits):
    v = _f(value)
    return round(v, ndigits) if v is not None else None


# --------------------------------------------------------------------- 汇总结果页
def _write_summary_en(wb, chambers, merged, cond, t_end):
    """按模板写「汇总结果」页的三张表，返回 (ws, pooled_title_row)。"""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet(en_labels.SUMMARY_SHEET_EN)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ---------- 表① 腔室级汇总 ----------
    ws.append(en_labels.SUMMARY_HEADERS)
    for j, spec in enumerate(en_richtext.SUMMARY_HEADER_SPECS, start=1):
        en_richtext.set_rt(ws.cell(row=1, column=j), spec, bold=True)
    ws.row_dimensions[1].height = 44.0          # 与模板一致
    for c in ws[1]:
        c.font = bold
        c.alignment = center
    for ch in chambers:
        vals = derived_params(ch["params"], t_end)
        # 模板列序：ID, μmax, Krel, Tlag, Tprolif, Tplat, Td, R²
        ws.append([ch["cb_id"], vals[0], vals[1], vals[6], vals[7], vals[8],
                   vals[5], vals[4]])

    # ---------- ② 培养条件块 ----------
    r = ws.max_row + 2
    ws.cell(row=r, column=1,
            value=en_labels.COND_BLOCK_TITLE.format(coded=cond["coded_text"])).font = \
        Font(bold=True, sz=12)
    for j, spec in enumerate(en_richtext.COND_BLOCK_SPECS, start=1):
        c = ws.cell(row=r + 1, column=j)
        en_richtext.set_rt(c, spec, size=10)    # 模板该行非粗体、sz 10
        c.font = Font(sz=10)
        c.alignment = center
    for j, text in enumerate(cond["values"], start=1):
        ws.cell(row=r + 2, column=j, value=text).alignment = center

    # ---------- ③ Pooled fitting results ----------
    r = ws.max_row + 2
    pooled_row = r
    ws.cell(row=r, column=1, value=en_labels.POOLED_TITLE).font = bold
    ws.cell(row=r, column=2, value=en_labels.POOLED_VALUE_HEADER).font = bold
    if merged:
        vals = derived_params(merged, t_end)
        # 模板行序：worksheets, points, μmax, Krel, Nrel,fit(0), λ, R²,
        #           Tlag, Tprolif, Tplat, t90
        values = [merged.get("n_worksheets"), merged.get("n_points"),
                  vals[0], vals[1], vals[2], vals[3], vals[4],
                  vals[6], vals[7], vals[8], vals[9]]
        for i, (label, value) in enumerate(zip(en_labels.POOLED_ROWS, values)):
            c = ws.cell(row=r + 1 + i, column=1)
            spec = en_richtext.POOLED_ROW_SPECS[i]
            if spec:
                en_richtext.set_rt(c, spec)
            else:
                c.value = label
            ws.cell(row=r + 1 + i, column=2, value=value)

    # ---------- 列宽（参照模板） ----------
    ws.column_dimensions["A"].width = 35.3
    for letter, width in zip("BCDEFGH", (28.4, 37.8, 21.2, 32.8, 22.9, 25.1, 15.7)):
        ws.column_dimensions[letter].width = width
    return ws, pooled_row


def _add_merged_plots_en(ws, chambers, merged, cond, row, logger, tag):
    """汇总结果页右侧嵌 3 张合并图，顶边对齐 Pooled fitting results 标题行。"""
    from openpyxl.utils import column_index_from_string

    if not merged:
        return
    series = []
    for ch in chambers:
        counts = np.asarray([v for v in ch["counts"] if v is not None], dtype=float)
        areas = np.asarray([v for v in ch["areas"] if v is not None], dtype=float)
        if counts.size >= 3:
            series.append((counts, areas))
    if not series:
        return

    n_area = config.PLOT_AREA_POINTS
    merged_t = np.concatenate([np.arange(c.size, dtype=float) for c, _ in series])
    merged_counts = np.concatenate([c for c, _ in series])
    t_area = np.concatenate([np.arange(min(c.size, n_area), dtype=float)
                             for c, _ in series])
    area_all = np.concatenate([a[:n_area] for _, a in series])
    avg_all = np.concatenate([relative_average_area(a[:n_area], c[:n_area])
                              for c, a in series])
    n_chambers = int(merged.get("n_worksheets") or len(series))
    cond_label = cond["title_label"]

    row = row + config.PLOT_MERGED_ROW_OFFSET
    start = column_index_from_string(config.PLOT_MERGED_START_COLUMN)
    _prepare_plot_band(ws, start, start + 24)
    cols = _merged_anchors(ws, row, start, 3)
    try:
        insert_image(ws, en_plotting.merged_growth_figure_en(
            merged_t, merged_counts, merged, n_chambers, cond_label),
            cols[0], config.PLOT_MERGED_SIZE)
        total_buf, avg_buf = en_plotting.merged_area_figures_en(
            t_area, area_all, avg_all, n_chambers, cond_label)
        insert_image(ws, total_buf, cols[1], config.PLOT_MERGED_SIZE)
        insert_image(ws, avg_buf, cols[2], config.PLOT_MERGED_SIZE)
    except Exception as exc:                                    # noqa: BLE001
        if logger:
            logger.warning("%s 汇总结果页图生成失败：%s", tag, exc)


# --------------------------------------------------------------------- 腔室页
def _write_chamber_en(wb, ch, merged, cond, logger, tag, plots=True):
    """CB-<k> 页：A:F 数据 + 10 行参数 + 4 张图（H2/H24/R2/R24）。"""
    from openpyxl.styles import Font

    ws = wb.create_sheet(ch["cb_id"])
    ws.append(en_labels.CHAMBER_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)

    param_vals = derived_params(ch["params"], ch["t_end"])
    for i, row in enumerate(ch["data_rows"], start=1):
        # row = (status, count, area)；A 列图片名按新腔室名重建
        status, count, area = row
        status = en_labels.STATUS_EN.get(str(status), status)
        new_row = [f"CB{ch['index']}_H{i}.tif", status, count, area, None, None]
        if i <= len(en_labels.CHAMBER_PARAM_LABELS):
            new_row[4] = en_labels.CHAMBER_PARAM_LABELS[i - 1]
            new_row[5] = param_vals[i - 1]
        ws.append(new_row)
        if i <= len(en_labels.CHAMBER_PARAM_LABELS):
            en_richtext.set_rt(ws.cell(row=ws.max_row, column=5),
                               en_richtext.CHAMBER_PARAM_SPECS[i - 1])

    for letter, width in zip("ABCDEF", (19.0, 15.4, 13.5, 23.5, 40.8, 15.1)):
        ws.column_dimensions[letter].width = width

    if not plots:
        return

    # ---------------- 4 张图 ----------------
    counts = np.asarray([r[1] for r in ch["data_rows"] if r[1] is not None],
                        dtype=float)
    areas = np.asarray([r[2] for r in ch["data_rows"] if r[2] is not None],
                       dtype=float)
    if counts.size < 3:
        if logger:
            logger.warning("%s %s 数据点不足（%d），跳过绘图",
                           tag, ch["cb_id"], counts.size)
        return
    t = np.arange(counts.size, dtype=float)
    anchors = config.PLOT_CHAMBER_ANCHORS
    cond_label = cond["title_label"]

    try:
        insert_image(ws, en_plotting.chamber_fit_figure_en(
            t, counts, ch["params"], ch["cb_id"], cond_label),
            anchors[0], config.PLOT_CHAMBER_SIZE)
    except Exception as exc:                                    # noqa: BLE001
        if logger:
            logger.warning("%s %s 拟合图失败：%s", tag, ch["cb_id"], exc)
    if merged:
        try:
            insert_image(ws, en_plotting.chamber_vs_merged_figure_en(
                t, counts, merged, ch["cb_id"], cond_label),
                anchors[1], config.PLOT_CHAMBER_SIZE)
        except Exception as exc:                                # noqa: BLE001
            if logger:
                logger.warning("%s %s 合并对比图失败：%s", tag, ch["cb_id"], exc)
    n = min(config.PLOT_AREA_POINTS, areas.size, counts.size)
    if n >= 3:
        try:
            total_buf, avg_buf = en_plotting.chamber_area_figures_en(
                t[:n], areas[:n],
                relative_average_area(areas[:n], counts[:n]),
                ch["cb_id"], cond_label)
            insert_image(ws, total_buf, anchors[2], config.PLOT_CHAMBER_SIZE)
            insert_image(ws, avg_buf, anchors[3], config.PLOT_CHAMBER_SIZE)
        except Exception as exc:                                # noqa: BLE001
            if logger:
                logger.warning("%s %s 面积图失败：%s", tag, ch["cb_id"], exc)


# --------------------------------------------------------------------- 入口
def _unify_fonts(wb):
    """所有非空单元格字体 → Times New Roman（保留粗体/字号/颜色，对齐模板）。"""
    from openpyxl.styles import Font

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                f = c.font
                if f is not None and f.name == "Times New Roman":
                    continue
                c.font = Font(
                    name="Times New Roman",
                    size=f.sz if f else None,
                    bold=f.b if f else False,
                    italic=f.i if f else False,
                    color=f.color if f else None,
                )


def build_english_workbook(out_path, chambers, merged, cond, logger=None, tag="",
                           plots=True):
    """
    chambers : [{index, cb_id, data_rows[(status,count,area)], counts, areas,
                 params{mu,K,N0,lam,r2,td,tplat,t90}, t_end}]
    merged   : {n_worksheets, n_points, mu, K, N0, lam, r2, tplat, t90} 或 None
    cond     : en_labels.condition_of_folder 的返回
    返回写入路径。
    """
    import openpyxl

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    t_end = max((ch["t_end"] for ch in chambers), default=None)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_summary, pooled_row = _write_summary_en(wb, chambers, merged, cond, t_end)
    for ch in chambers:
        _write_chamber_en(wb, ch, merged, cond, logger, tag, plots=plots)
    if plots:
        _add_merged_plots_en(ws_summary, chambers, merged, cond, pooled_row,
                             logger, tag)
    _unify_fonts(wb)
    wb.save(out_path)

    if logger:
        logger.info("  %s 写出 %s（%d 个腔室页签，共 %d 张图）",
                    tag, out_path.name, len(chambers),
                    len(chambers) * 4 + 3 if plots else 0)
    return out_path

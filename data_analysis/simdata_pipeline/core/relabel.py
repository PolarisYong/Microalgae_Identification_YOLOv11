# -*- coding: utf-8 -*-
"""
按新条件改写数据组文件夹的全部标签
==================================
把一个数据组文件夹（工作簿 + PPT + 文件名）里所有「条件标识」与「原始腔室溯源」
统一替换成指定条件。难点在于：工作簿内嵌图与 PPT 配图都是 **PNG 位图**，
图标题是画在图上的像素，改不了字，只能按同一套绘图代码重绘。因此本模块：

    1. 从原工作簿的腔室页签**读回时间序列**（数据一字不改，不重新处理 tif）
    2. 交给 core.writer 重新生成工作簿：腔室页签统一为 CB1..CBn、图标题带新条件、
       「数据来源」页签整体删除、腔室页 A 列的图片名去掉原始流道前缀
    3. 用 python-pptx 原地替换 PPT 里的配图（保留原版式与位置），并把
       CH<k>-CB<j> 标签改写为 CB<n>

「去掉原始溯源」的规则（与需求一致）：
    CH11_CB1_H1.tif   → CB1_H1.tif        （腔室页 A 列）
    CH1_CB1           → CB1                （页签名 / 汇总表腔室名称）
    CH11-CB1          → CB1                （PPT 图片标签）
    数据来源           → 整页删除
"""
import re
from pathlib import Path

import config

# 汇总表「腔室名称」列里的键，如 CH1_CB1
CHAMBER_KEY_RE = re.compile(r"^CH\d+_CB(\d+)$")
# 腔室页 A 列图片名结尾的小时序号，如 _H1.tif / _H96.tif
HOUR_SUFFIX_RE = re.compile(r"(_H\d+\.tif)$", re.IGNORECASE)
# 腔室页 A 列图片名的原始流道前缀，如 CH11_CB1_
SOURCE_PREFIX_RE = re.compile(r"^CH\d+_CB\d+_")
# PPT 上的腔室标签，如 CH11-CB1
PPT_LABEL_RE = re.compile(r"^CH\d+-CB(\d+)$")

SUMMARY_SHEET = "汇总"
SOURCE_SHEET = "数据来源"


# --------------------------------------------------------------------- 读回腔室数据
def summary_chamber_order(wb):
    """从「汇总」页第一张表读出腔室顺序（即 workbook 里腔室页签的顺序）。"""
    if SUMMARY_SHEET not in wb.sheetnames:
        return []
    order = []
    for row in wb[SUMMARY_SHEET].iter_rows(min_row=2, values_only=True):
        value = row[0] if row else None
        if isinstance(value, str) and CHAMBER_KEY_RE.fullmatch(value.strip()):
            order.append(value.strip())
        elif order:
            break
    return order


def _suffix_label(value, if_missing=""):
    """CH11_CB1_H1.tif → CB1_H1.tif；无流道前缀时退化为补前缀。"""
    name = str(value)
    hour = HOUR_SUFFIX_RE.search(name)
    if hour:
        return f"{if_missing}{hour.group(1)}"
    if SOURCE_PREFIX_RE.match(name):
        return SOURCE_PREFIX_RE.sub(f"{if_missing}_", name)
    return name


def _relabel_row(row, label):
    """腔室页一行 → A 列去掉原始流道溯源后的新行。"""
    new_row = [v for v in row]
    if new_row and isinstance(new_row[0], str):
        new_row[0] = _suffix_label(new_row[0], label)
    return new_row


def read_source(xlsx_path, condition, logger=None, tag=""):
    """
    读回原工作簿 → (chambers, original_summary_rows)

    chambers : writer.build_workbook 需要的腔室字典列表，页签名统一为 CB1..CBn，
               plot_name 已按新条件拼好（图标题直接用它）。
    original_summary_rows : 「汇总」页原样快照，供改写后逐格比对。
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        order = summary_chamber_order(wb)
        if not order:
            order = [s for s in wb.sheetnames
                     if s not in (SUMMARY_SHEET, SOURCE_SHEET)
                     and CHAMBER_KEY_RE.fullmatch(s)]
            if logger:
                logger.warning("%s 「汇总」页未识别到腔室列表，改用页签顺序（%d 个）",
                               tag, len(order))
        if not order:
            raise ValueError(f"{Path(xlsx_path).name} 里没有找到任何腔室页签")

        original_summary = []
        if SUMMARY_SHEET in wb.sheetnames:
            original_summary = [[v for v in r]
                                for r in wb[SUMMARY_SHEET].iter_rows(values_only=True)]
        return extract_chambers(wb, order, condition, logger, tag), original_summary
    finally:
        wb.close()


def extract_chambers(wb, order, condition, logger=None, tag=""):
    """腔室页签 → chambers 列表（new_name=CBn、plot_name=条件_CBn、A 列去溯源）。"""
    chambers = []
    for index, sheet in enumerate(order, 1):
        label = f"CB{index}"
        ws = wb[sheet]
        rows = [[v for v in r] for r in ws.iter_rows(values_only=True)]
        header, body = rows[0], rows[1:]

        counts = [r[config.COUNTS_COLUMN_INDEX] for r in body
                  if len(r) > config.COUNTS_COLUMN_INDEX
                  and r[config.COUNTS_COLUMN_INDEX] is not None]
        areas = [r[config.AREA_COLUMN_INDEX] for r in body
                 if len(r) > config.AREA_COLUMN_INDEX
                 and r[config.AREA_COLUMN_INDEX] is not None]
        params = {}
        for r in body:
            if len(r) > 5 and r[4] is not None:
                key = str(r[4]).strip()
                if key and key != "参数名":
                    params[key] = r[5]

        chambers.append({
            "new_name": label,
            "plot_name": f"{condition}_{label}",     # 图标题里的腔室名
            "rows": [header] + [_relabel_row(r, label) for r in body],
            "counts": counts,
            "areas": areas,
            "params": params,
        })
        if logger:
            logger.info("%s %-10s ← %-10s  %d 个数据点 / %d 个参数",
                        tag, label, sheet, len(counts), len(params))
    return chambers


# --------------------------------------------------------------------- 重建 PPT
def _shape_sort_key(shape):
    return (shape.top, shape.left)


def rebuild_pptx(src_pptx, dst_pptx, pngs, logger=None, tag=""):
    """
    复制原 PPT 并原地换图 + 改写腔室标签。

    pngs 顺序必须与幻灯片里按 (top, left) 排序后的图片顺序一致（即腔室顺序）。
    返回 (图片数, 改写标签数)。
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(src_pptx))
    if not prs.slides:
        raise ValueError(f"{Path(src_pptx).name} 里没有幻灯片")
    slide = prs.slides[0]

    pictures = sorted([sh for sh in slide.shapes
                       if sh.shape_type == MSO_SHAPE_TYPE.PICTURE], key=_shape_sort_key)
    labels = sorted([sh for sh in slide.shapes
                     if sh.has_text_frame and PPT_LABEL_RE.fullmatch(sh.text_frame.text.strip())],
                    key=_shape_sort_key)
    if len(pictures) != len(pngs):
        raise ValueError(f"PPT 里有 {len(pictures)} 张图，但待替换的 PNG 有 {len(pngs)} 张，"
                         f"数量不一致，已中止以免错配")
    if labels and len(labels) != len(pictures):
        if logger:
            logger.warning("%s PPT 标签数 %d ≠ 图片数 %d，仅按几何位置配对",
                           tag, len(labels), len(pictures))

    # ① 换图：直接替换图片部件的二进制内容，位置 / 尺寸 / 版式完全不动
    for pic, data in zip(pictures, pngs):
        r_id = pic._element.blipFill.blip.rEmbed       # noqa: SLF001
        part = pic.part.related_part(r_id)
        part._blob = data                               # noqa: SLF001
        if hasattr(part, "_image"):                     # 清掉 python-pptx 的缓存
            del part._image

    # ② 改标签：每张图配它正上方、同一列的文本框（第 k 张图 → CBk）
    relabeled = 0
    for index, pic in enumerate(pictures, 1):
        target, best = None, None
        for lab in labels:
            if lab.left != pic.left:
                continue
            distance = pic.top - lab.top
            if distance <= 0:
                continue
            if best is None or distance < best:
                target, best = lab, distance
        if target is None:
            if logger:
                logger.warning("%s 第 %d 张图没找到对应标签框", tag, index)
            continue
        _set_text(target, f"CB{index}")
        relabeled += 1

    prs.save(str(dst_pptx))
    return len(pictures), relabeled


def _set_text(shape, text):
    """保留首个 run 的字体格式，替换整段文字。"""
    tf = shape.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        para.text = text


# --------------------------------------------------------------------- 校验
def collect_cell_texts(xlsx_path, skip_sheets=()):
    """收集工作簿里所有单元格的文本，用于检查是否还残留原始溯源。"""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        out = []
        for name in wb.sheetnames:
            if name in skip_sheets:
                continue
            for row in wb[name].iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        out.append((name, v))
        return out
    finally:
        wb.close()


def find_leftovers(texts, patterns):
    """在单元格文本里找仍匹配给定正则的项。"""
    hits = []
    for sheet, text in texts:
        for pat in patterns:
            if pat.search(text):
                hits.append((sheet, text))
                break
    return hits

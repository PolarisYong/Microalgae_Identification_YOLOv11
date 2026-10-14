# -*- coding: utf-8 -*-
"""
幻灯片抽取
==========
python-pptx 没有「跨演示文稿复制页」的官方 API。为了保证内容 100% 保真
（图、表、字体、版式全保留），这里的做法是：

    以源 pptx 为底稿打开 → 只删除不需要的页 → 另存为新文件

比「新建空演示 + 复制 shape」可靠得多，因为后者要自己复制图片、
图表等关联部件，容易丢内容。

页序说明：growth_curves pptx 的第 i 页 ≡ selected__groups_<i>_<n>.xlsx，
页标题形如「数据组 3（n = 18）」。本模块会解析标题做一次交叉校验，
不一致时记 WARNING（仍按组序号取页）。
"""
import re
from pathlib import Path

# 标题形如：数据组 3（n = 18）
TITLE_RE = re.compile(r"数据组\s*(\d+).*?n\s*=\s*(\d+)")


def slide_titles(pptx_path):
    """[(页码, 标题文本)]，取每页第一个非空文本框。"""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        text = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip().splitlines()[0]
                break
        out.append((i, text))
    return out


def check_title(page_no, title, group_index, group_size):
    """
    校验第 page_no 页的标题与 (组序号, 腔室数) 是否吻合。
    返回 None 表示通过，否则返回说明文字。
    """
    m = TITLE_RE.search(title or "")
    if not m:
        return None                      # 标题格式变了，跳过校验
    t_idx, t_size = int(m.group(1)), int(m.group(2))
    if t_idx != group_index or t_size != group_size:
        return (f"第 {page_no} 页标题为「{title}」，与组 {group_index}"
                f"（n = {group_size}）不符")
    return None


def extract_slides(src_pptx, keep_pages, out_path):
    """
    从 src_pptx 抽取 keep_pages（1 起的原页码）另存为 out_path。

    返回 (实际保留的页码列表, 越界/缺失的页码列表, 源总页数)。
    保留页按原顺序排列。
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn

    src_pptx = Path(src_pptx)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(src_pptx))
    total = len(prs.slides)
    wanted = sorted({int(p) for p in keep_pages})
    keep = [p for p in wanted if 1 <= p <= total]
    missing = [p for p in wanted if not (1 <= p <= total)]

    sld_id_lst = prs.slides._sldIdLst
    for i in range(total - 1, -1, -1):          # 倒序删除，避免下标错位
        if (i + 1) not in keep:
            sld_id = sld_id_lst[i]
            r_id = sld_id.get(qn("r:id"))
            prs.part.drop_rel(r_id)
            sld_id_lst.remove(sld_id)

    prs.save(str(out_path))
    return keep, missing, total

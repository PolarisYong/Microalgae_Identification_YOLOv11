# -*- coding: utf-8 -*-
"""
selected__groups_*.xlsx 英文化
==============================
复制源工作簿，把 selected_N / strict_M 页签的表头按
en_labels.SELECTED_HEADER_MAP 翻成英文；summary 页与全部数据原样保留
（metric/value 本身已是英文，腔室名称属于原始溯源，不改）。
"""
from pathlib import Path

from core import en_labels, en_richtext


def translate_selected(src_path, dst_path, logger=None, tag=""):
    """返回翻译的表头单元格数。"""
    import openpyxl

    wb = openpyxl.load_workbook(src_path)
    translated = 0
    for name in wb.sheetnames:
        ws = wb[name]
        if not (name.startswith("selected_") or name.startswith("strict_")):
            continue
        for cell in ws[1]:
            if isinstance(cell.value, str):
                new = en_labels.SELECTED_HEADER_MAP.get(cell.value.strip())
                if new:
                    cell.value = new
                    spec = en_richtext.SELECTED_HEADER_SPECS.get(new)
                    if spec:
                        en_richtext.set_rt(
                            cell, spec,
                            bold=bool(cell.font.b),
                            size=cell.font.sz or 11)
                    translated += 1
    dst_path = Path(dst_path)
    dst_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst_path)
    wb.close()
    if logger:
        logger.info("  %s selected 表头翻译 %d 处 → %s", tag, translated,
                    dst_path.name)
    return translated
